#!/usr/bin/env python3
"""Contrôle qualité d'une présentation au regard de la charte Triskell.

Usage :
    python check_deck.py presentation.pptx [--json]

Fonctionne sur n'importe quel .pptx (généré par build_deck.py ou fait à la main).
Contrôles :
  1. Polices hors charte (autre que Montserrat / Arial / police du thème)
  2. Couleurs de texte et de forme hors palette (tolérance de proximité)
  3. Consignes du modèle oubliées (« A MODIFIER », « Insérer… », « NOM CLIENT », « xxx »)
  4. Champs encore à compléter (« [à compléter] »)
  5. Titres absents, trop longs, ou deux slides avec le même titre
  6. Densité : trop de mots sur une slide, trop de puces dans un bloc
  7. Plus d'un élément orange « d'attention » par slide (hors intertitres)
  8. Séquence : couverture en premier, fin en dernier, séparateur toutes les 6 slides au plus
Code de sortie : 0 si aucun bloquant, 1 sinon.
"""
import json
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
T = json.load(open(os.path.join(ROOT, "tokens", "tokens.json"), encoding="utf-8"))
PALETTE = [RGBColor.from_string(v.lstrip("#")) for v in T["couleurs"].values()]
FONTS_OK = {T["polices"][k] for k in ("titre", "texte", "display", "repli")} | {"Montserrat Medium", "Montserrat Bold", None, "+mj-lt", "+mn-lt"}
FORBIDDEN = [s.lower() for s in T["consignes_interdites"]] + ["xxx"]
MAX_WORDS = 140
LAYOUT_NAMES = {"cover": "2_Disposition personnalisée", "section": "Transition-vert"}


def dist(a, b):
    return sum((int(str(a)[i:i + 2], 16) - int(str(b)[i:i + 2], 16)) ** 2 for i in (0, 2, 4)) ** 0.5


def off_palette(rgb):
    return min(dist(rgb, p) for p in PALETTE) > 40


def check(path):
    prs = Presentation(path)
    issues = []  # (gravité, slide, message)
    titles = {}
    seq = []
    since_section = 0
    for n, slide in enumerate(prs.slides, 1):
        lname = slide.slide_layout.name
        kind = "cover" if lname == LAYOUT_NAMES["cover"] else "section" if lname == LAYOUT_NAMES["section"] else \
            "closing" if "MERCI" in " ".join(s.text_frame.text for s in slide.slide_layout.shapes if s.has_text_frame) else \
            "content" if lname.startswith("Contenu") else "free"
        seq.append(kind)
        since_section = 0 if kind in ("section", "cover") else since_section + 1
        if since_section == T["regles"]["slides_max_sans_separateur"] + 1:
            issues.append(("conseil", n, f"plus de {T['regles']['slides_max_sans_separateur']} slides de contenu sans séparateur de section"))
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        if kind == "content":
            if not title and slide.shapes.title is not None:
                issues.append(("important", n, "slide sans titre"))
            elif len(title) > T["regles"]["titre_max_caracteres"]:
                issues.append(("conseil", n, f"titre de {len(title)} caractères : un titre = un message court"))
            if title:
                if title in titles:
                    issues.append(("conseil", n, f"même titre que la slide {titles[title]} : « {title[:50]} »"))
                titles.setdefault(title, n)
        words, oranges = 0, 0
        for shp in slide.shapes:
            try:
                if shp.fill.type == 1 and off_palette(shp.fill.fore_color.rgb):
                    issues.append(("important", n, f"forme « {shp.name} » de couleur hors charte #{shp.fill.fore_color.rgb}"))
            except Exception:
                pass
            frames = [shp.text_frame] if shp.has_text_frame else []
            if getattr(shp, "has_table", False) and shp.has_table:
                frames += [c.text_frame for r in shp.table.rows for c in r.cells]
            for tf in frames:
                text = tf.text
                low = text.lower()
                for f in FORBIDDEN:
                    if re.search(r"(?<!\w)" + re.escape(f) + r"(?!\w)", low):
                        issues.append(("bloquant", n, f"consigne ou texte de modèle oublié : « {f} »"))
                if T["texte_manquant"] in text:
                    issues.append(("bloquant", n, f"champ encore « {T['texte_manquant']} » : {' / '.join(text.split(chr(10)))[:70]}"))
                words += len(text.split())
                bullets = sum(1 for p in tf.paragraphs if p.text.strip().startswith(("▪", "•", "-")))
                if bullets > T["regles"]["puces_max_par_bloc"]:
                    issues.append(("conseil", n, f"{bullets} puces dans un même bloc"))
                for p in tf.paragraphs:
                    for r in p.runs:
                        if r.font.name not in FONTS_OK:
                            issues.append(("important", n, f"police hors charte « {r.font.name} » ({r.text[:30]})"))
                        try:
                            rgb = r.font.color.rgb if r.font.color and r.font.color.type == 1 else None
                        except Exception:
                            rgb = None
                        if rgb is not None:
                            if off_palette(rgb):
                                issues.append(("important", n, f"couleur de texte hors charte #{rgb} ({r.text[:30]})"))
                            big = r.font.size and r.font.size.pt >= 30
                            if dist(rgb, RGBColor.from_string("FD7F01")) < 40 and big:
                                oranges += 1
        if words > MAX_WORDS and kind == "content":
            issues.append(("conseil", n, f"{words} mots : slide trop dense (cible < {MAX_WORDS})"))
        if oranges > 1 and kind == "content":
            issues.append(("conseil", n, f"{oranges} éléments orange de grande taille : un seul point d'attention par slide"))
    if seq and seq[0] != "cover":
        issues.append(("important", 1, "la présentation ne commence pas par la couverture officielle"))
    if seq and seq[-1] != "closing":
        issues.append(("conseil", len(seq), "pas de slide « MERCI ! » en fin de présentation"))
    # dédoublonnage
    seen, out = set(), []
    for i in issues:
        if i not in seen:
            seen.add(i)
            out.append(i)
    return out


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    res = check(sys.argv[1])
    order = {"bloquant": 0, "important": 1, "conseil": 2}
    res.sort(key=lambda x: (order[x[0]], x[1]))
    if "--json" in sys.argv:
        print(json.dumps([{"gravite": g, "slide": s, "message": m} for g, s, m in res], ensure_ascii=False, indent=2))
    else:
        counts = {g: sum(1 for r in res if r[0] == g) for g in order}
        print(f"Contrôle charte Triskell : {counts['bloquant']} bloquant(s), {counts['important']} important(s), {counts['conseil']} conseil(s)")
        for g, s, m in res:
            print(f"  [{g.upper()}] slide {s} : {m}")
    sys.exit(1 if any(r[0] == "bloquant" for r in res) else 0)
