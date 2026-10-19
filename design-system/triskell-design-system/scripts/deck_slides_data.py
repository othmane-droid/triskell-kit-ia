"""Slides de données et de projet : graphique, tableau, phases, jalons, feuille de route, étapes,
avant / après, matrice, check-list, actions, personnes, équipe, offres, image + texte."""
import os

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Cm, Pt

from deck_core import (C, FONT, FONT_BLACK, MISSING, SERIES, STATE, TOKENS, ZONE, asset, check_forbidden, disc, is_empty,
                       line, line_h, lines_needed, note_missing, panel, rect, shape, style_run, textbox, warn)
from deck_slides_base import BOTTOM, content, notes, source

CHART_TYPES = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
               "stacked": XL_CHART_TYPE.COLUMN_STACKED, "line": XL_CHART_TYPE.LINE_MARKERS, "donut": XL_CHART_TYPE.DOUGHNUT}


def _font(tf_or_font, size, color="texte_secondaire", bold=False):
    f = tf_or_font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = C[color]


def s_chart(prs, s):
    """Graphique natif PowerPoint (modifiable) dans la palette, avec un message à retenir à droite."""
    slide, y = content(prs, s)
    kind = s.get("chart", "column")
    data = CategoryChartData()
    cats = [c if not is_empty(c) else f"[catégorie {i + 1}]" for i, c in enumerate(s["categories"])]
    if cats != s["categories"]:
        note_missing("catégories du graphique")
    data.categories = cats
    for se in s["series"]:
        if any(v is None for v in se["values"]):
            note_missing(f"valeur de la série {se.get('name') or 'sans nom'}")
        data.add_series(se.get("name") or "Série", [v if v is not None else 0 for v in se["values"]])
    has_side = bool(s.get("takeaway"))
    w = 21.5 if has_side else ZONE["w"]
    h = BOTTOM - y - 0.9
    gf = slide.shapes.add_chart(CHART_TYPES[kind], Cm(ZONE["x"]), Cm(y), Cm(w), Cm(h), data)
    ch = gf.chart
    ch.font.name, ch.font.size = FONT, Pt(12)
    ch.font.color.rgb = C["texte_secondaire"]
    multi = len(s["series"]) > 1 or kind == "donut"
    ch.has_legend = multi
    if multi:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        _font(ch.legend.font, 12)
    plot = ch.plots[0]
    if kind == "donut":
        pts = plot.series[0].points
        for i in range(len(cats)):
            pts[i].format.fill.solid()
            pts[i].format.fill.fore_color.rgb = SERIES[i % len(SERIES)]
    else:
        if kind in ("column", "bar", "stacked"):
            plot.gap_width = 70
            plot.overlap = 100 if kind == "stacked" else -10
        hi = s.get("highlight_category")
        for i, se in enumerate(plot.series):
            fill = se.format.fill if kind != "line" else None
            if kind == "line":
                se.format.line.color.rgb = SERIES[i % len(SERIES)]
                se.format.line.width = Pt(2.75)
                se.smooth = False
            else:
                fill.solid()
                fill.fore_color.rgb = SERIES[i % len(SERIES)]
                if hi is not None and len(plot.series) == 1:
                    for j, cat in enumerate(cats):
                        if cat == hi:
                            pt = se.points[j]
                            pt.format.fill.solid()
                            pt.format.fill.fore_color.rgb = C["orange"]
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = C["trait"]
        va.format.line.fill.background()
        _font(va.tick_labels.font, 11, "texte_attenue")
        ca = ch.category_axis
        ca.format.line.color.rgb = C["trait"]
        _font(ca.tick_labels.font, 12)
    if s.get("labels", True):
        plot.has_data_labels = True
        dl = plot.data_labels
        _font(dl.font, 12, "texte", True)
        dl.number_format = s.get("number_format", '# ##0')
        dl.number_format_is_linked = False
        if kind in ("column", "bar"):
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
    if has_side:
        x = ZONE["x"] + w + 0.8
        sw = ZONE["w"] - w - 0.8
        rect(slide, x, y, sw, h, "panneau", rounded=True)
        rect(slide, x, y, sw, 0.2, "orange")
        paras = []
        if s.get("key_figure"):
            paras.append({"text": s["key_figure"], "size": 46, "bold": True, "font": FONT_BLACK, "color": "orange", "space_after": 4})
        paras.append({"text": s["takeaway"], "size": 18, "bold": True, "space_after": 10})
        if s.get("takeaway_detail"):
            paras.append({"text": s["takeaway_detail"], "size": 14, "color": "texte_secondaire"})
        textbox(slide, x + 0.6, y + 0.8, sw - 1.2, h - 1.2, paras, min_size=11)
    source(slide, s, required=True)
    notes(slide, s)


def s_table(prs, s):
    slide, y = content(prs, s)
    cols, rows = s["columns"], s["rows"]
    if len(rows) > 12:
        warn(f"{len(rows)} lignes : découper en (1/2), (2/2)")
    h_avail = BOTTOM - y - 0.9
    row_h = min(1.35, h_avail / (len(rows) + 1))
    size = 15 if row_h >= 1.3 else (13 if row_h >= 1.0 else 11)
    gf = slide.shapes.add_table(len(rows) + 1, len(cols), Cm(ZONE["x"]), Cm(y), Cm(ZONE["w"]), Cm(row_h * (len(rows) + 1)))
    tbl = gf.table
    widths = s.get("col_widths") or [1] * len(cols)
    for j, wv in enumerate(widths):
        tbl.columns[j].width = Cm(ZONE["w"] * wv / sum(widths))
    header = s.get("header_color", "orange_tableau")
    hi_row = s.get("highlight_row")

    def cell(c, text, bold, color, fill, label, align_right=False):
        italic = False
        if is_empty(text):
            note_missing(label)
            text, color, italic = MISSING, "orange", True
        check_forbidden(str(text))
        c.text = ""
        p = c.text_frame.paragraphs[0]
        if align_right:
            from pptx.enum.text import PP_ALIGN
            p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = str(text)
        style_run(r, size, color, bold, italic)
        c.fill.solid()
        c.fill.fore_color.rgb = C[fill]
        c.margin_left = c.margin_right = Cm(0.3)
        c.margin_top = c.margin_bottom = Cm(0.1)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE

    num_cols = set(s.get("numeric_columns", []))
    for j, h in enumerate(cols):
        cell(tbl.cell(0, j), h, True, "blanc", header, "en-tête", j in num_cols)
    for i, row in enumerate(rows, 1):
        is_hi = hi_row is not None and i - 1 == hi_row
        is_total = s.get("total_row") and i == len(rows)
        for j in range(len(cols)):
            v = row[j] if j < len(row) else None
            fill = "fond_orange" if is_hi else ("fond_doux" if is_total else ("gris_fond" if i % 2 == 0 else "blanc"))
            cell(tbl.cell(i, j), v, j == 0 or is_total, "texte", fill, f"{row[0] if row else ''} / {cols[j]}", j in num_cols)
    source(slide, s)
    notes(slide, s)


def s_phases(prs, s):
    slide, y = content(prs, s)
    ph = s["phases"]
    w = ZONE["w"] / len(ph)
    for i, p in enumerate(ph):
        cx = ZONE["x"] + i * w
        cur = p.get("current")
        shp = shape(slide, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, cx, y, w + 0.45, 1.7,
                    "orange" if cur else ("vert" if i % 2 == 0 else "vert_clair"))
        tf = shp.text_frame
        tf.margin_left = Cm(1.0 if i else 0.5)
        r = tf.paragraphs[0].add_run()
        r.text = p["name"]
        style_run(r, 17, "blanc", True)
        rect(slide, cx + 0.2, y + 2.2, w - 0.6, BOTTOM - y - 2.6, "fond_orange" if cur else "panneau", rounded=True)
        textbox(slide, cx + 0.7, y + 2.6, w - 1.5, 1.0, [{"text": p.get("dates"), "size": 15, "bold": True, "color": "orange" if cur else "vert", "label": f"dates {p['name']}"}])
        paras = [{"text": it, "size": 15, "bullet": True, "bullet_color": "orange" if cur else "vert", "space_after": 8} for it in p.get("items", [])]
        if p.get("deliverable"):
            paras.append({"prefix": "Livrable : ", "text": p["deliverable"], "size": 14, "color": "texte_secondaire"})
        textbox(slide, cx + 0.7, y + 3.7, w - 1.5, BOTTOM - y - 4.3, paras, min_size=10)
    notes(slide, s)


def s_timeline(prs, s):
    """Jalons datés sur une ligne : en haut / en bas en alternance."""
    slide, y = content(prs, s)
    ms = s["milestones"]
    n = len(ms)
    mid = y + (BOTTOM - y) / 2
    x0, x1 = ZONE["x"] + 0.8, ZONE["x"] + ZONE["w"] - 0.8
    rect(slide, x0, mid - 0.08, x1 - x0, 0.16, "trait")
    done_until = max([i for i, m in enumerate(ms) if m.get("done")] or [-1])
    if done_until >= 0:
        rect(slide, x0, mid - 0.08, (x1 - x0) * done_until / max(1, n - 1), 0.16, "vert")
    w = (x1 - x0) / max(1, n - 1) if n > 1 else 10
    for i, m in enumerate(ms):
        cx = x0 + i * w if n > 1 else (x0 + x1) / 2
        color = "orange" if m.get("current") else ("vert" if m.get("done") else "texte_attenue")
        d = 1.0 if m.get("current") else 0.7
        disc(slide, cx - d / 2, mid - d / 2, d, color)
        up = i % 2 == 0
        bw = min(7.5, w * 1.6) if n > 1 else 10
        by = mid - 5.0 if up else mid + 1.1
        bx = min(max(cx - bw / 2, ZONE["x"]), ZONE["x"] + ZONE["w"] - bw)
        line(slide, cx, mid - (0.5 if up else -0.5), cx, mid - (1.0 if up else -1.1), "trait", 1)
        textbox(slide, bx, by, bw, 3.9,
                [{"text": m.get("date"), "size": 15, "bold": True, "color": color if color != "texte_attenue" else "vert", "align": "center", "label": "date du jalon", "space_after": 3},
                 {"text": m["label"], "size": 16, "bold": True, "align": "center", "space_after": 3}] +
                ([{"text": m["detail"], "size": 13, "color": "texte_secondaire", "align": "center"}] if m.get("detail") else []),
                anchor="bottom" if up else "top", min_size=10, align="left" if bx == ZONE["x"] else ("right" if bx + bw >= ZONE["x"] + ZONE["w"] - 0.01 else "center"))
    notes(slide, s)


def s_roadmap(prs, s):
    """Feuille de route : chantiers × périodes, barres et jalons."""
    slide, y = content(prs, s)
    periods, lanes = s["periods"], s["lanes"]
    lab_w = 6.8
    x0 = ZONE["x"] + lab_w
    pw = (ZONE["w"] - lab_w) / len(periods)
    for j, p in enumerate(periods):
        rect(slide, x0 + j * pw + 0.05, y, pw - 0.1, 1.0, "vert" if j == s.get("today") else "panneau")
        textbox(slide, x0 + j * pw, y, pw, 1.0, [{"text": p, "size": 13, "bold": True, "color": "blanc" if j == s.get("today") else "texte", "align": "center"}], anchor="middle", quiet=True)
    lane_h = min(2.2, (BOTTOM - y - 2.2) / max(1, len(lanes)))
    for i, ln in enumerate(lanes):
        ly = y + 1.4 + i * lane_h
        if i % 2 == 0:
            rect(slide, ZONE["x"], ly, ZONE["w"], lane_h, "gris_fond")
        textbox(slide, ZONE["x"] + 0.3, ly, lab_w - 0.5, lane_h, [{"text": ln["name"], "size": 15, "bold": True}], anchor="middle", min_size=10)
        for b in ln.get("bars", []):
            bx = x0 + b["start"] * pw + 0.1
            bw = (b["end"] - b["start"]) * pw - 0.2
            bar = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bx, ly + lane_h * 0.22, bw, lane_h * 0.56, "orange" if b.get("highlight") else ("vert" if i % 2 == 0 else "vert_clair"), radius=0.3)
            tf = bar.text_frame
            tf.margin_left = Cm(0.3)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.word_wrap = True
            if b.get("label"):
                r = tf.paragraphs[0].add_run()
                r.text = b["label"]
                style_run(r, 12 if lines_needed(b["label"], 12, bw - 0.4) == 1 else 10, "blanc", True)
    for m in s.get("milestones", []):
        mx = x0 + m["at"] * pw
        line(slide, mx, y + 1.2, mx, y + 1.4 + len(lanes) * lane_h, "orange", 1.5)
        dia = shape(slide, MSO_SHAPE.DIAMOND, mx - 0.35, y + 1.45 + len(lanes) * lane_h, 0.7, 0.7, "orange")
        textbox(slide, mx - 3.5, y + 2.25 + len(lanes) * lane_h, 7, 1.0, [{"text": m["label"], "size": 12, "bold": True, "color": "orange", "align": "center"}], quiet=True)
    notes(slide, s)


def s_steps(prs, s):
    """Démarche en étapes numérotées reliées."""
    slide, y = content(prs, s)
    st = s["steps"]
    n = len(st)
    w = ZONE["w"] / n
    cy = y + 0.4
    line(slide, ZONE["x"] + w / 2, cy + 0.9, ZONE["x"] + ZONE["w"] - w / 2, cy + 0.9, "trait", 2)
    for i, p in enumerate(st):
        cx = ZONE["x"] + i * w + w / 2
        hi = p.get("highlight")
        disc(slide, cx - 0.9, cy, 1.8, "orange" if hi else "vert", i + 1, size=20)
        paras = [{"text": p["title"], "size": 18, "bold": True, "align": "center", "space_after": 8}]
        if p.get("text"):
            paras.append({"text": p["text"], "size": 14, "color": "texte_secondaire", "align": "center", "space_after": 8})
        if p.get("duration"):
            paras.append({"text": p["duration"], "size": 13, "bold": True, "color": "orange" if hi else "vert", "align": "center"})
        textbox(slide, cx - w / 2 + 0.3, cy + 2.4, w - 0.6, BOTTOM - cy - 2.8, paras, min_size=10)
    notes(slide, s)


def s_before_after(prs, s):
    slide, y = content(prs, s)
    b, a = s["before"], s["after"]
    gap = 2.6
    w = (ZONE["w"] - gap) / 2
    h = BOTTOM - y - 0.3
    rect(slide, ZONE["x"], y, w, h, "gris_fond", rounded=True)
    textbox(slide, ZONE["x"] + 0.7, y + 0.6, w - 1.4, h - 1.0,
            [{"text": b.get("label", "AUJOURD'HUI").upper(), "size": 12, "bold": True, "color": "texte_attenue", "space_after": 6},
             {"text": b["title"], "size": 21, "bold": True, "color": "texte_secondaire", "space_after": 12}] +
            [{"text": it, "size": 15, "bullet": True, "bullet_color": "texte_attenue", "color": "texte_secondaire", "space_after": 8} for it in b["items"]], min_size=11)
    ax = ZONE["x"] + w + 0.5
    shape(slide, MSO_SHAPE.RIGHT_ARROW, ax, y + h / 2 - 0.9, gap - 1.0, 1.8, "orange")
    x2 = ZONE["x"] + w + gap
    rect(slide, x2, y, w, h, "fond_doux", rounded=True)
    rect(slide, x2, y, w, 0.2, "vert")
    textbox(slide, x2 + 0.7, y + 0.6, w - 1.4, h - 1.0,
            [{"text": a.get("label", "AVEC TRISKELL").upper(), "size": 12, "bold": True, "color": "vert", "space_after": 6},
             {"text": a["title"], "size": 21, "bold": True, "space_after": 12}] +
            [{"text": it, "size": 15, "bullet": True, "bullet_color": "vert", "space_after": 8} for it in a["items"]], min_size=11)
    notes(slide, s)


def s_matrix(prs, s):
    """Matrice 2 × 2 (priorisation valeur / effort, risques…). Quadrants : haut-gauche, haut-droit, bas-gauche, bas-droit."""
    slide, y = content(prs, s)
    mh = BOTTOM - y - 1.3
    mw = 17.5 if s.get("reading") else ZONE["w"] - 1.6
    x0 = ZONE["x"] + 1.4
    qs = s.get("quadrants", ["", "", "", ""])
    hq = s.get("highlight_quadrant", 0)
    hw, hh = mw / 2, mh / 2
    for k in range(4):
        qx, qy = x0 + (k % 2) * hw, y + (k // 2) * hh
        rect(slide, qx + 0.05, qy + 0.05, hw - 0.1, hh - 0.1, "fond_doux" if k == hq else ("panneau" if k in (1, 2) else "gris_fond"))
        textbox(slide, qx + 0.35, qy + 0.25, hw - 0.7, 0.9, [{"text": qs[k].upper(), "size": 12, "bold": True, "color": "vert" if k == hq else "texte_attenue"}], quiet=True)
    textbox(slide, x0, y + mh + 0.15, mw, 0.8, [{"text": s.get("x_label", "") + "  →", "size": 13, "bold": True, "color": "texte_secondaire", "align": "center"}], quiet=True)
    lab = slide.shapes.add_textbox(Cm(ZONE["x"] + 0.5 - mh / 2), Cm(y + mh / 2 - 0.4), Cm(mh), Cm(0.8))
    lab.rotation = 270
    r = lab.text_frame.paragraphs[0].add_run()
    r.text = s.get("y_label", "") + "  →"
    style_run(r, 13, "texte_secondaire", True)
    from pptx.enum.text import PP_ALIGN
    lab.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    for it in s["items"]:
        px, py = x0 + it["x"] * mw, y + (1 - it["y"]) * mh
        disc(slide, px - 0.35, py - 0.35, 0.7, "orange" if it.get("highlight") else "vert")
        lw = min(7.5, hw - 0.8)
        right = it["x"] > 0.55
        textbox(slide, px - 0.6 - lw if right else px + 0.6, py - 0.45, lw, 0.9,
                [{"text": it["label"], "size": 14, "bold": True, "align": "right" if right else "left"}], quiet=True)
    if s.get("reading"):
        rx = x0 + mw + 1.0
        panel(slide, rx, y, ZONE["x"] + ZONE["w"] - rx, mh, s.get("reading_title", "Lecture"),
              [{"text": t, "size": 15, "bullet": True, "space_after": 10} for t in s["reading"]], accent="orange")
    notes(slide, s)


STATUS = {"oui": "icone-oui.png", "non": "icone-non.png"}


def s_checklist(prs, s):
    slide, y = content(prs, s)
    items = s["items"]
    show_owner = any(it.get("owner") for it in items)
    h_avail = BOTTOM - y - (1.9 if s.get("verdict", True) else 0.3)
    row = min(1.55, h_avail / (len(items) + 1))
    qw = 17.0 if show_owner else 19.0
    x0 = ZONE["x"]
    hdr = [("Point de contrôle", x0, qw, "left"), ("Oui / Non", x0 + qw, 2.4, "center")]
    if show_owner:
        hdr.append(("Porteur", x0 + qw + 2.6, 2.4, "left"))
    cx_comment = x0 + qw + (5.2 if show_owner else 2.6)
    hdr.append(("Commentaire", cx_comment, 8, "left"))
    for t, hx, hw, al in hdr:
        textbox(slide, hx, y, hw, row, [{"text": t, "size": 13, "bold": True, "color": "texte_attenue", "align": al}], quiet=True)
    qsize = 16
    while qsize > 11 and any(lines_needed(it.get("question") or MISSING, qsize, qw - 0.4) * line_h(qsize) + 0.15 > row for it in items):
        qsize -= 0.5
    oui = non = vide = 0
    for i, it in enumerate(items, 1):
        cy = y + i * row
        if i % 2 == 0:
            rect(slide, x0 - 0.2, cy - 0.05, ZONE["w"] + 0.4, row, "gris_fond")
        textbox(slide, x0, cy, qw - 0.4, row, [{"text": it["question"], "size": qsize, "bold": it.get("critical", False)}], anchor="middle", min_size=qsize)
        st = (it.get("status") or "").lower()
        ic = 0.85
        if st in STATUS:
            slide.shapes.add_picture(asset(STATUS[st]), Cm(x0 + qw + 0.8), Cm(cy + (row - ic) / 2), Cm(ic), Cm(ic))
            oui += st == "oui"
            non += st == "non"
        else:
            vide += 1
            note_missing("statut : " + str(it.get("question") or "point de contrôle"))
            textbox(slide, x0 + qw, cy, 2.4, row, [{"text": "?", "size": 17, "bold": True, "color": "orange", "align": "center"}], anchor="middle")
        if show_owner:
            textbox(slide, x0 + qw + 2.6, cy, 2.4, row, [{"text": it.get("owner") or "—", "size": 15, "bold": True, "color": "vert"}], anchor="middle")
        textbox(slide, cx_comment, cy, ZONE["x"] + ZONE["w"] - cx_comment, row, [{"text": it.get("comment") or "", "size": 14, "color": "texte_secondaire", "allow_empty": True}], anchor="middle", min_size=9)
    if s.get("verdict", True):
        critical_ko = [it for it in items if it.get("critical") and (it.get("status") or "").lower() != "oui"]
        if non == 0 and vide == 0:
            label, color = "GO", "vert"
        elif critical_ko:
            label, color = "NO GO", "rouge"
        else:
            label, color = "GO sous réserve", "orange"
        txt = f"Verdict : {label}  ·  {oui} oui · {non} non · {vide} à vérifier"
        if critical_ko:
            txt += f"  ·  bloquant : {critical_ko[0]['question'][:60]}"
        textbox(slide, x0, BOTTOM - 1.5, ZONE["w"] - 3.6, 1.4, [{"text": txt, "size": 15, "bold": True, "color": "blanc"}], fill=color, anchor="middle")
    notes(slide, s)


def s_actions(prs, s):
    rows = []
    for team in s["teams"]:
        for a in team["actions"]:
            a = a if isinstance(a, dict) else {"action": a}
            rows.append([team["team"], a.get("action"), a.get("owner"), a.get("due")])
    s_table(prs, {"title": s["title"], "subtitle": s.get("subtitle"), "columns": ["Équipe", "Action", "Porteur", "Échéance"],
                  "rows": rows, "col_widths": [2.2, 7, 2.2, 1.8], "header_color": "vert", "notes": s.get("notes")})


def s_people(prs, s):
    from deck_slides_base import s_fields
    cols = [{"heading": g["heading"], "fields": [{"label": p["role"], "value": p.get("name")} for p in g["people"]]} for g in s["groups"]]
    s_fields(prs, {"title": s["title"], "subtitle": s.get("subtitle"), "columns": cols, "notes": s.get("notes")})


def s_team(prs, s):
    """Équipe projet : une carte par personne, initiales dans un disque."""
    slide, y = content(prs, s)
    ppl = s["members"]
    per_row = min(4, len(ppl))
    rows = -(-len(ppl) // per_row)
    gap = 0.8
    w = (ZONE["w"] - (per_row - 1) * gap) / per_row
    h = min(5.6, (BOTTOM - y - 0.3 - (rows - 1) * gap) / rows)
    for i, p in enumerate(ppl):
        cx = ZONE["x"] + (i % per_row) * (w + gap)
        cy = y + (i // per_row) * (h + gap)
        rect(slide, cx, cy, w, h, "panneau", rounded=True)
        initials = "".join(part[0] for part in (p.get("name") or "?").split()[:2]).upper()
        disc(slide, cx + 0.6, cy + 0.6, 1.9, "orange" if p.get("lead") else "vert", initials, size=18)
        textbox(slide, cx + 2.9, cy + 0.5, w - 3.3, 2.2,
                [{"text": p.get("name"), "size": 17, "bold": True, "label": "nom", "space_after": 2},
                 {"text": p.get("role"), "size": 14, "bold": True, "color": "vert", "label": "rôle"}], anchor="middle", min_size=10)
        if p.get("text"):
            textbox(slide, cx + 0.6, cy + 2.9, w - 1.2, h - 3.2, [{"text": p["text"], "size": 14, "color": "texte_secondaire"}], min_size=10)
    notes(slide, s)


def s_offer(prs, s):
    """Formules ou options chiffrées, la recommandée mise en avant."""
    slide, y = content(prs, s)
    offers = s["offers"]
    n = len(offers)
    gap = 0.8
    w = (ZONE["w"] - (n - 1) * gap) / n
    h = min(BOTTOM - y - 0.3, 11.0)
    for i, o in enumerate(offers):
        cx = ZONE["x"] + i * (w + gap)
        rec = o.get("recommended")
        rect(slide, cx, y, w, h, "fond_orange" if rec else "panneau", rounded=True)
        rect(slide, cx, y, w, 0.22, "orange" if rec else "vert")
        if rec:
            textbox(slide, cx + w - 4.4, y + 0.6, 3.8, 0.85, [{"text": o.get("badge", "Recommandé"), "size": 11, "bold": True, "color": "blanc", "align": "center", "space_after": 0}], fill="orange", anchor="middle", pad=0.08)
        textbox(slide, cx + 0.7, y + 0.6, w - 5.4 if rec else w - 1.4, 1.2, [{"text": o["name"], "size": 20, "bold": True}], min_size=13)
        textbox(slide, cx + 0.7, y + 1.9, w - 1.4, 2.2,
                [{"text": o.get("price"), "size": 34, "bold": True, "font": FONT_BLACK, "color": "orange" if rec else "texte_titre", "label": "prix", "space_after": 0},
                 {"text": o.get("period", ""), "size": 13, "color": "texte_secondaire", "allow_empty": True}], min_size=18)
        line(slide, cx + 0.7, y + 4.4, cx + w - 0.7, y + 4.4, "trait", 1)
        paras = ([{"text": o["pitch"], "size": 14, "color": "texte_secondaire", "space_after": 10}] if o.get("pitch") else []) + \
                [{"text": it, "size": 14, "bullet": True, "bullet_color": "orange" if rec else "vert", "space_after": 6} for it in o.get("items", [])]
        textbox(slide, cx + 0.7, y + 4.8, w - 1.4, h - 5.2, paras, min_size=10)
    if s.get("footnote"):
        textbox(slide, ZONE["x"], BOTTOM - 0.1, ZONE["w"] - 5, 0.6, [{"text": s["footnote"], "size": 10, "color": "texte_attenue"}], quiet=True)
    notes(slide, s)


def s_image_text(prs, s):
    """Photo d'un côté, message et preuves de l'autre. Image : chemin fourni ou photo du modèle."""
    slide, y = content(prs, s)
    img = s.get("image", "reunion-equipe")
    path = img if os.path.exists(str(img)) else asset("photos", f"{img}.jpg")
    left = s.get("side", "left") == "left"
    iw = 14.5
    ih = BOTTOM - y - 0.3
    ix = ZONE["x"] if left else ZONE["x"] + ZONE["w"] - iw
    if os.path.exists(path):
        pic = slide.shapes.add_picture(path, Cm(ix), Cm(y), Cm(iw), Cm(ih))
        # recadrage pour garder les proportions
        from PIL import Image
        W, H = Image.open(path).size
        target = iw / ih
        if W / H > target:
            cut = (1 - target * H / W) / 2
            pic.crop_left = pic.crop_right = cut
        else:
            cut = (1 - (W / H) / target) / 2
            pic.crop_top = pic.crop_bottom = cut
    else:
        warn(f"image introuvable : {img}")
    tx = ZONE["x"] + iw + 1.2 if left else ZONE["x"]
    tw = ZONE["w"] - iw - 1.2
    paras = []
    if s.get("kicker"):
        paras.append({"text": s["kicker"].upper(), "size": 12, "bold": True, "color": "vert", "space_after": 8})
    paras.append({"text": s["heading"], "size": 26, "bold": True, "color": "texte_titre", "space_after": 14})
    if s.get("text"):
        paras.append({"text": s["text"], "size": 16, "color": "texte_secondaire", "space_after": 12})
    paras += [{"text": it, "size": 16, "bullet": True, "space_after": 8} for it in s.get("items", [])]
    textbox(slide, tx, y + 0.3, tw, ih - 0.6, paras, anchor="middle", min_size=11)
    notes(slide, s)


TYPES = {"chart": s_chart, "table": s_table, "phases": s_phases, "timeline": s_timeline, "roadmap": s_roadmap,
         "steps": s_steps, "before_after": s_before_after, "matrix": s_matrix, "checklist": s_checklist,
         "actions": s_actions, "people": s_people, "team": s_team, "offer": s_offer, "image_text": s_image_text}
