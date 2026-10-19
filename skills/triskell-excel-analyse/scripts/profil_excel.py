"""Profil d'un classeur Excel : feuilles, colonnes, types, anomalies.

Usage : python profil_excel.py fichier.xlsx [--json profil.json]
Ne modifie jamais le fichier. Affiche un rapport en français (Markdown).
"""
import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from collections import Counter

from openpyxl import load_workbook

DATE_TXT = re.compile(r"^\s*\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}\s*$")
NUM_TXT = re.compile(r"^\s*-?\d[\d\s ]*([.,]\d+)?\s*(€|%)?\s*$")


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def find_header(ws, max_scan=15):
    """Première ligne dont au moins 2 cellules sont des textes et 60 % sont remplies."""
    best = 1
    for r in range(1, min(ws.max_row, max_scan) + 1):
        vals = [c.value for c in ws[r]]
        filled = [v for v in vals if v not in (None, "")]
        texts = [v for v in filled if isinstance(v, str)]
        if len(texts) >= 2 and len(filled) >= 0.6 * max(1, ws.max_column):
            return r
    return best


def kind(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return "vide"
    if isinstance(v, bool):
        return "booléen"
    if isinstance(v, (int, float)):
        return "nombre"
    if isinstance(v, (dt.date, dt.datetime)):
        return "date"
    if isinstance(v, str) and DATE_TXT.match(v):
        return "date en texte"
    if isinstance(v, str) and NUM_TXT.match(v):
        return "nombre en texte"
    return "texte"


def profile_column(name, cells, n_rows):
    vals = [c.value for c in cells]
    kinds = Counter(kind(v) for v in vals)
    main = max((k for k in kinds if k != "vide"), key=lambda k: kinds[k], default="vide")
    col = {"colonne": name, "type_principal": main, "types": dict(kinds),
           "vides": kinds.get("vide", 0), "anomalies": []}
    anomalies = col["anomalies"]
    for k in ("date en texte", "nombre en texte"):
        if kinds.get(k) and main in ("date", "nombre"):
            rows = [c.row for c, v in zip(cells, vals) if kind(v) == k]
            anomalies.append(f"{kinds[k]} valeur(s) « {k} » : lignes {rows[:10]}")
    if kinds.get("vide") and kinds["vide"] < n_rows:
        rows = [c.row for c, v in zip(cells, vals) if kind(v) == "vide"]
        anomalies.append(f"{kinds['vide']} cellule(s) vide(s) : lignes {rows[:10]}")
    texts = [v for v in vals if isinstance(v, str) and v.strip()]
    if texts:
        groups = {}
        for v in texts:
            groups.setdefault(norm(v), set()).add(v)
        variants = [sorted(s) for s in groups.values() if len(s) > 1]
        for s in variants[:10]:
            anomalies.append(f"variantes d'écriture d'une même valeur : {s}")
        spaces = [c.row for c, v in zip(cells, vals) if isinstance(v, str) and v != v.strip()]
        if spaces:
            anomalies.append(f"espaces en début / fin : lignes {spaces[:10]}")
        col["valeurs_distinctes"] = len(set(texts))
        if len(set(texts)) <= 12:
            col["valeurs"] = dict(Counter(texts).most_common())
    nums = [(c.row, float(v)) for c, v in zip(cells, vals) if kind(v) == "nombre"]
    if nums:
        xs = sorted(x for _, x in nums)
        med = xs[len(xs) // 2]
        col.update(min=xs[0], max=xs[-1], somme=round(sum(xs), 4), mediane=med)
        neg = [r for r, x in nums if x < 0]
        if neg and xs[len(xs) // 2] > 0:
            anomalies.append(f"valeur(s) négative(s) : lignes {neg[:10]}")
        if med > 0:
            out = [(r, x) for r, x in nums if x > 5 * med]
            if out:
                anomalies.append(f"valeur(s) > 5 × la médiane ({med}) : {out[:10]}")
    dates = [(c.row, v) for c, v in zip(cells, vals) if isinstance(v, (dt.date, dt.datetime))]
    if dates:
        ds = [v for _, v in dates]
        col.update(date_min=min(ds).strftime("%d/%m/%Y"), date_max=max(ds).strftime("%d/%m/%Y"))
        today = dt.datetime.now()
        fut = [r for r, v in dates if (v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)) > today]
        if fut and len(fut) < len(dates):  # colonne mixte passé / futur (une colonne d'échéances est toute future)
            anomalies.append(f"{len(fut)} date(s) postérieure(s) à aujourd'hui ({today:%d/%m/%Y}) : lignes {fut[:10]} "
                             "(prévisionnel ou erreur de saisie ? à vérifier)")
    return col


def profile_sheet(ws, ws_formulas):
    h = find_header(ws)
    headers = [c.value for c in ws[h]]
    n_rows = max(0, ws.max_row - h)
    rows = [tuple(c.value for c in r) for r in ws.iter_rows(min_row=h + 1)]
    rows = [r for r in rows if any(v not in (None, "") for v in r)]
    dup = [h + 1 + i for i, r in enumerate(rows) if r in rows[:i]]
    formulas = sum(1 for r in ws_formulas.iter_rows() for c in r
                   if isinstance(c.value, str) and c.value.startswith("="))
    sheet = {"feuille": ws.title, "visible": ws.sheet_state == "visible", "ligne_entete": h,
             "lignes_donnees": len(rows), "colonnes": len(headers), "formules": formulas,
             "cellules_fusionnees": len(ws.merged_cells.ranges), "doublons_exacts": dup,
             "details": []}
    if len(rows) == 0 or ws.max_column < 2:
        return sheet
    for j, name in enumerate(headers, start=1):
        cells = [ws.cell(r, j) for r in range(h + 1, ws.max_row + 1)]
        sheet["details"].append(profile_column(name or f"(colonne {j} sans titre)", cells, n_rows))
    return sheet


def to_markdown(path, sheets):
    out = [f"# Profil du classeur : {path}", ""]
    for s in sheets:
        vis = "" if s["visible"] else " (feuille masquée)"
        out.append(f"## Feuille « {s['feuille']} »{vis}")
        out.append(f"En-tête ligne {s['ligne_entete']} · {s['lignes_donnees']} lignes de données · "
                   f"{s['colonnes']} colonnes · {s['formules']} formule(s) · "
                   f"{s['cellules_fusionnees']} plage(s) fusionnée(s)")
        if s["doublons_exacts"]:
            out.append(f"- **Doublons exacts** : lignes {s['doublons_exacts']}")
        if not s["details"]:
            out.append("- Pas de tableau de données détecté (feuille de texte).")
        else:
            out += ["", "| Colonne | Type principal | Vides | Résumé |", "|---|---|---|---|"]
            for c in s["details"]:
                res = []
                if "somme" in c:
                    res.append(f"min {c['min']} · max {c['max']} · somme {c['somme']}")
                if "date_min" in c:
                    res.append(f"du {c['date_min']} au {c['date_max']}")
                if "valeurs" in c:
                    res.append(", ".join(f"{k} ({v})" for k, v in c["valeurs"].items()))
                elif "valeurs_distinctes" in c:
                    res.append(f"{c['valeurs_distinctes']} valeurs distinctes")
                out.append(f"| {c['colonne']} | {c['type_principal']} | {c['vides']} | {' ; '.join(res)} |")
            anos = [(c["colonne"], a) for c in s["details"] for a in c["anomalies"]]
            if anos:
                out += ["", "**Anomalies détectées (rien n'a été modifié)**"]
                out += [f"- {col} : {a}" for col, a in anos]
        out.append("")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("fichier")
    ap.add_argument("--json")
    a = ap.parse_args()
    try:
        wb = load_workbook(a.fichier, data_only=True)
        wbf = load_workbook(a.fichier, data_only=False)
    except Exception as e:  # fichier protégé, .xls ancien, corrompu
        sys.exit(f"Impossible d'ouvrir le fichier : {e}. Format .xlsx attendu (enregistrer un .xls en .xlsx).")
    sheets = [profile_sheet(wb[n], wbf[n]) for n in wb.sheetnames]
    print(to_markdown(a.fichier, sheets))
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump(sheets, f, ensure_ascii=False, indent=2, default=str)


if __name__ == "__main__":
    main()
