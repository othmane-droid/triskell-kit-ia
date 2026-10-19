"""Copie propre d'un classeur : feuilles d'origine intactes + « Synthèse » + « Données nettoyées »
+ « Journal des corrections », avec graphiques Excel natifs.

Usage : python copie_propre.py source.xlsx spec.json sortie.xlsx
La spec (JSON) décrit les décisions de nettoyage validées et les tableaux de synthèse :
voir references/copie-propre-spec.md. Le fichier source n'est jamais modifié.
"""
import argparse
import datetime as dt
import json
import re
import sys
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VERT, ORANGE, TEXTE = "00897B", "FD7F01", "3F3F3F"
ENTETE = dict(fill=PatternFill("solid", fgColor=VERT), font=Font(bold=True, color="FFFFFF"))
TITRE = Font(bold=True, size=14, color=TEXTE)
SOUS_TITRE = Font(bold=True, size=11, color=ORANGE)
AGGS = {"somme": sum, "sum": sum, "nombre": len, "count": len,
        "moyenne": lambda xs: sum(xs) / len(xs) if xs else 0, "mean": lambda xs: sum(xs) / len(xs) if xs else 0,
        "max": max, "min": min}


def lire_table(ws, h: int):
    headers = [c.value for c in ws[h]]
    rows = []
    for r in ws.iter_rows(min_row=h + 1):
        vals = [c.value for c in r]
        if any(v not in (None, "") for v in vals):
            rows.append({"_ligne": r[0].row, "_controle": [], **dict(zip(headers, vals))})
    return headers, rows


def en_nombre(v):
    if isinstance(v, str):
        s = v.replace(" ", "").replace(" ", "").replace("€", "").replace("%", "").replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return v
    return v


def en_date(v, fmt):
    if isinstance(v, str):
        try:
            return dt.datetime.strptime(v.strip(), fmt)
        except ValueError:
            return v
    return v


def condition_vraie(v, cond: str, ref: dt.datetime) -> bool:
    if cond == "vide":
        return v is None or (isinstance(v, str) and not v.strip())
    if cond in ("futur", "passé"):
        if not isinstance(v, (dt.date, dt.datetime)):
            return False
        v = v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)
        return v > ref if cond == "futur" else v < ref
    m = re.match(r"^\s*(<=|>=|<|>|==|!=)\s*(-?[\d.]+)\s*$", cond)
    if not m or not isinstance(v, (int, float)):
        return False
    op, x = m.group(1), float(m.group(2))
    return {"<": v < x, ">": v > x, "<=": v <= x, ">=": v >= x, "==": v == x, "!=": v != x}[op]


def appliquer(rows, decisions, journal, ref: dt.datetime):
    """Applique les décisions dans l'ordre ; chaque modification est journalisée."""
    for d in decisions:
        did, col, act = d.get("id", "?"), d.get("colonne"), d["action"]
        if act == "supprimer_doublons_exacts":
            vus, gardes = set(), []
            for r in rows:
                cle = tuple((k, str(v)) for k, v in r.items() if not k.startswith("_"))
                if cle in vus:
                    journal.append((did, r["_ligne"], "(ligne entière)", "doublon", "retirée", d.get("motif", "doublon exact")))
                else:
                    vus.add(cle)
                    gardes.append(r)
            rows[:] = gardes
            continue
        for r in rows:
            avant = r.get(col)
            apres = avant
            if act == "supprimer_espaces" and isinstance(avant, str):
                apres = re.sub(r"\s+", " ", avant).strip()
            elif act == "remplacer" and avant in d["valeurs"]:
                apres = d["valeurs"][avant]
            elif act == "texte_en_nombre":
                apres = en_nombre(avant)
            elif act == "texte_en_date":
                apres = en_date(avant, d.get("format", "%d/%m/%Y"))
            elif act == "signaler" and condition_vraie(avant, d["condition"], ref):
                r["_controle"].append(f"{did} : {d.get('message', d['condition'])}")
                continue
            if apres != avant:
                r[col] = apres
                journal.append((did, r["_ligne"], col, avant, apres, d.get("motif", act)))
    return rows


def ecrire_entete(ws, row: int, values):
    for j, v in enumerate(values, start=1):
        c = ws.cell(row, j, v)
        c.fill, c.font = ENTETE["fill"], ENTETE["font"]
        c.alignment = Alignment(wrap_text=True, vertical="center")


def format_cellule(c):
    if isinstance(c.value, (dt.date, dt.datetime)):
        c.number_format = "DD/MM/YYYY"
    elif isinstance(c.value, float):
        c.number_format = "#,##0.00"


def feuille_donnees(wb, headers, rows):
    ws = wb.create_sheet("Données nettoyées")
    ecrire_entete(ws, 1, ["Ligne d'origine"] + headers + ["Contrôle"])
    for i, r in enumerate(rows, start=2):
        vals = [r["_ligne"]] + [r.get(h) for h in headers] + ["; ".join(r["_controle"]) or "OK"]
        for j, v in enumerate(vals, start=1):
            format_cellule(ws.cell(i, j, v))
        if r["_controle"]:
            ws.cell(i, len(vals)).font = Font(bold=True, color=ORANGE)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    largeurs(ws)


def feuille_journal(wb, journal, decisions):
    ws = wb.create_sheet("Journal des corrections")
    ecrire_entete(ws, 1, ["Décision", "Ligne d'origine", "Colonne", "Avant", "Après", "Motif"])
    for i, e in enumerate(journal, start=2):
        for j, v in enumerate(e, start=1):
            format_cellule(ws.cell(i, j, v if not isinstance(v, list) else str(v)))
    n = len(journal) + 3
    ws.cell(n, 1, "Décisions appliquées").font = SOUS_TITRE
    ecrire_entete(ws, n + 1, ["Décision", "Action", "Colonne", "Détail"])
    for k, d in enumerate(decisions, start=n + 2):
        detail = d.get("motif") or d.get("message") or json.dumps(d.get("valeurs", d.get("condition", "")), ensure_ascii=False)
        for j, v in enumerate([d.get("id"), d["action"], d.get("colonne", "(toutes)"), detail], start=1):
            ws.cell(k, j, v)
    largeurs(ws)


def pivot(rows, s):
    groupes = defaultdict(list)
    for r in rows:
        if s.get("exclure_signalees") and r["_controle"]:
            continue
        v = r.get(s["valeur"]) if s.get("valeur") else 1
        if isinstance(v, (int, float)):
            groupes[r.get(s["index"]) or "(vide)"].append(v)
    f = AGGS[s.get("agg", "somme")]
    lignes = sorted(([k, round(f(v), 2)] for k, v in groupes.items()), key=lambda x: -x[1])
    return [s["index"], s.get("libelle_valeur", s.get("valeur", "Nombre"))], lignes


def graphique(ws, s, r0: int, nb: int, ncol: int, ancre: str):
    kind = s.get("graphique")
    if not kind or nb == 0:
        return
    ch = {"barres": BarChart, "lignes": LineChart, "camembert": PieChart}[kind]()
    if kind == "barres":
        ch.type = "bar" if s.get("horizontal") else "col"
    ch.title = s["titre"]
    series = s.get("series") or list(range(2, ncol + 1))
    for col in series:
        ch.add_data(Reference(ws, min_col=col, min_row=r0, max_row=r0 + nb), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + nb))
    if kind != "camembert":
        for i, ser in enumerate(ch.series):
            ser.graphicalProperties.solidFill = [VERT, ORANGE, TEXTE][i % 3]
            ser.graphicalProperties.line.solidFill = [VERT, ORANGE, TEXTE][i % 3]
    ch.height, ch.width = 7.5, 14
    ws.add_chart(ch, ancre)


def feuille_synthese(wb, spec, rows):
    ws = wb.create_sheet("Synthèse", 0)
    ws["A1"] = spec.get("titre", "Synthèse")
    ws["A1"].font = TITRE
    ws["A2"] = spec.get("sous_titre", "Copie de travail : les feuilles d'origine sont inchangées.")
    r = 4
    blocs = [("Chiffres clés", "indicateurs"), ("Hypothèses de calcul", "hypotheses"), ("Ce qu'on en conclut", "conclusions")]
    for titre, cle in blocs[:1]:
        if spec.get(cle):
            ws.cell(r, 1, titre).font = SOUS_TITRE
            ecrire_entete(ws, r + 1, ["Indicateur", "Valeur", "Source / calcul"])
            for k, ind in enumerate(spec[cle], start=r + 2):
                for j, v in enumerate([ind["libelle"], ind["valeur"], ind.get("source", "")], start=1):
                    format_cellule(ws.cell(k, j, v))
            r = r + 3 + len(spec[cle])
    for s in spec.get("syntheses", []):
        if s.get("type", "pivot") == "pivot":
            entetes, lignes = pivot(rows, s)
        else:
            entetes, lignes = s["colonnes"], s["lignes"]
        ws.cell(r, 1, s["titre"]).font = SOUS_TITRE
        ecrire_entete(ws, r + 1, entetes)
        for k, ligne in enumerate(lignes, start=r + 2):
            for j, v in enumerate(ligne, start=1):
                format_cellule(ws.cell(k, j, v))
        if s.get("note"):
            ws.cell(r + 2 + len(lignes), 1, s["note"]).font = Font(italic=True, color=TEXTE)
        graphique(ws, s, r + 1, len(lignes), len(entetes), f"{get_column_letter(len(entetes) + 2)}{r}")
        r += max(len(lignes) + 4, 17 if s.get("graphique") else 0)
    for titre, cle in blocs[1:]:
        if spec.get(cle):
            ws.cell(r, 1, titre).font = SOUS_TITRE
            for k, t in enumerate(spec[cle], start=r + 1):
                ws.cell(k, 1, f"• {t}")
            r += len(spec[cle]) + 2
    ws.column_dimensions["A"].width = 38
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18


def largeurs(ws):
    for col in ws.columns:
        n = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, n + 2), 45)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("spec")
    ap.add_argument("sortie")
    a = ap.parse_args()
    if a.source == a.sortie:
        sys.exit("La sortie doit être un nouveau fichier : le fichier source n'est jamais écrasé.")
    with open(a.spec, encoding="utf-8") as f:
        spec = json.load(f)
    wb = load_workbook(a.source)                     # formules et mise en forme d'origine conservées
    valeurs = load_workbook(a.source, data_only=True)
    if spec["feuille"] not in valeurs.sheetnames:
        sys.exit(f"Feuille « {spec['feuille']} » introuvable. Feuilles : {valeurs.sheetnames}")
    headers, rows = lire_table(valeurs[spec["feuille"]], spec.get("ligne_entete", 1))
    journal = []
    ref = dt.datetime.strptime(spec["date_reference"], "%d/%m/%Y") if spec.get("date_reference") else dt.datetime.now()
    rows = appliquer(rows, spec.get("decisions", []), journal, ref)
    feuille_synthese(wb, spec, rows)
    feuille_donnees(wb, headers, rows)
    feuille_journal(wb, journal, spec.get("decisions", []))
    wb.active = 0
    wb.save(a.sortie)
    signalees = sum(1 for r in rows if r["_controle"])
    print(f"OK : {a.sortie} · {len(rows)} lignes nettoyées · {len(journal)} modification(s) journalisée(s) · "
          f"{signalees} ligne(s) signalée(s) · feuilles : {wb.sheetnames}")


if __name__ == "__main__":
    main()
