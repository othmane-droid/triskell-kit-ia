"""Extrait les exigences numérotées d'un cahier des charges (.docx, .pdf, .txt, .md) avec leur page ou section.

Usage : python extraire_exigences.py cdc.docx [--motif REGEX] [--json exigences.json]
Motif par défaut : identifiants du type « EX-12 », « REQ 3.4 », « F-07 », « Exigence 12 ».
Le niveau (obligatoire / souhaité / optionnel) est repris s'il est écrit près de l'identifiant.
Ne remplace pas la lecture : les exigences non numérotées (phrases « doit », « devra ») sont listées à part.
"""
import argparse
import json
import re
import sys

MOTIF = r"^\s*((?:EX|REQ|EXG|F|NF|T|E)[- ]?\d+(?:\.\d+)*|Exigence\s+\d+(?:\.\d+)*)\s*(?:\(([^)]*)\))?\s*[:.\-–]\s*(.+)$"
NIVEAUX = {"obligatoire": "Obligatoire", "impératif": "Obligatoire", "imperatif": "Obligatoire",
           "souhaité": "Souhaité", "souhaite": "Souhaité", "optionnel": "Optionnel", "facultatif": "Optionnel"}
OBLIGATION = re.compile(r"\b(doit|doivent|devra|devront|est exigé|sont exigés|obligatoirement|impérativement)\b", re.I)


def lire_blocs(path: str):
    """Renvoie une liste de (emplacement, texte)."""
    if path.lower().endswith(".docx"):
        from docx import Document
        doc, section, blocs = Document(path), "début du document", []
        for p in doc.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            if p.style.name.lower().startswith(("heading", "titre")):
                section = t
            blocs.append((f"§ {section}", t))
        for i, table in enumerate(doc.tables, start=1):
            for r, row in enumerate(table.rows, start=1):
                blocs.append((f"tableau {i}, ligne {r}", " | ".join(c.text.strip() for c in row.cells)))
        return blocs
    if path.lower().endswith(".pdf"):
        from pypdf import PdfReader
        blocs = []
        for n, page in enumerate(PdfReader(path).pages, start=1):
            for line in (page.extract_text() or "").splitlines():
                if line.strip():
                    blocs.append((f"p. {n}", line.strip()))
        if not blocs:
            sys.exit("PDF sans texte extractible (scan ?) : demander une version texte ou faire un OCR.")
        return blocs
    with open(path, encoding="utf-8") as f:
        return [("ligne " + str(i), l.strip()) for i, l in enumerate(f, start=1) if l.strip()]


def extraire(blocs, motif: str):
    rx, exigences, autres = re.compile(motif, re.I), [], []
    for lieu, texte in blocs:
        m = rx.match(texte)
        if m:
            niveau_brut = (m.group(2) or "").strip().lower()
            niveau = next((v for k, v in NIVEAUX.items() if k in niveau_brut), "Non précisé")
            exigences.append({"id": m.group(1).replace(" ", "-"), "niveau": niveau,
                              "exigence": m.group(3).strip(), "emplacement": lieu})
        elif OBLIGATION.search(texte):
            autres.append({"emplacement": lieu, "texte": texte})
    return exigences, autres


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("fichier")
    ap.add_argument("--motif", default=MOTIF)
    ap.add_argument("--json")
    a = ap.parse_args()
    exigences, autres = extraire(lire_blocs(a.fichier), a.motif)
    ids = [e["id"] for e in exigences]
    doublons = sorted({i for i in ids if ids.count(i) > 1})
    print(f"{len(exigences)} exigence(s) numérotée(s) · "
          + ", ".join(f"{n} {k.lower()}" for k, n in sorted(
              {k: sum(1 for e in exigences if e['niveau'] == k) for k in {e['niveau'] for e in exigences}}.items())))
    if doublons:
        print(f"ATTENTION identifiants en double : {doublons}")
    for e in exigences:
        print(f"- {e['id']} [{e['niveau']}] ({e['emplacement']}) {e['exigence'][:90]}")
    if autres:
        print(f"\n{len(autres)} phrase(s) d'obligation non numérotée(s), à vérifier :")
        for o in autres:
            print(f"- ({o['emplacement']}) {o['texte'][:110]}")
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump({"exigences": exigences, "obligations_non_numerotees": autres}, f, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
