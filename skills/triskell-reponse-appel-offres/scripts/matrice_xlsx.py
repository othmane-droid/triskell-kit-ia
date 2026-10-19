"""Génère le classeur de réponse à un appel d'offres : Synthèse, Matrice de conformité,
Questions acheteur, Grille go / no go. Les compteurs et le score sont des formules Excel :
ils se mettent à jour quand l'équipe modifie un statut ou une note.

Usage : python matrice_xlsx.py analyse.json "<Acheteur> - <Référence> - matrice de conformité.xlsx"
Format du JSON : references/matrice-conformite.md (§ Format JSON).
"""
import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

VERT, ORANGE, TEXTE = "00897B", "FD7F01", "3F3F3F"
STATUTS = {"Couvert": "D5F0EC", "Partiel": "FFE7CC", "Non couvert": "F8D0CC", "À vérifier": "E6E6E6"}
FIN = Side(style="thin", color="BFBFBF")
BORD = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)
COLS = [("N°", "id", 9), ("Niveau", "niveau", 12), ("Catégorie", "categorie", 16), ("Exigence (texte du CDC)", "exigence", 48),
        ("Emplacement", "emplacement", 18), ("Réponse proposée", "reponse", 55), ("Statut", "statut", 13),
        ("Source Triskell", "source", 30), ("Pièce / preuve", "piece", 22), ("Responsable", "responsable", 16),
        ("Commentaire interne (à retirer avant envoi)", "commentaire_interne", 40)]


def entete(ws, row, values, widths=None):
    for j, v in enumerate(values, start=1):
        c = ws.cell(row, j, v)
        c.fill, c.font, c.border = PatternFill("solid", fgColor=VERT), Font(bold=True, color="FFFFFF"), BORD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[c.column_letter].width = widths[j - 1]


def matrice(wb, exigences):
    ws = wb.create_sheet("Matrice de conformité")
    entete(ws, 1, [c[0] for c in COLS], [c[2] for c in COLS])
    for i, e in enumerate(exigences, start=2):
        if e.get("statut") not in STATUTS:
            sys.exit(f"{e.get('id')} : statut « {e.get('statut')} » invalide. Valeurs : {list(STATUTS)}")
        for j, (_, key, _) in enumerate(COLS, start=1):
            c = ws.cell(i, j, e.get(key) or ("[à compléter]" if key == "reponse" else ""))
            c.alignment, c.border = Alignment(wrap_text=True, vertical="top"), BORD
            if key == "commentaire_interne" and c.value:
                c.font = Font(italic=True, color=ORANGE)
    last = len(exigences) + 1
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{last}")
    for statut, couleur in STATUTS.items():
        ws.conditional_formatting.add(f"A2:K{last}", FormulaRule(formula=[f'$G2="{statut}"'],
                                                               fill=PatternFill("solid", fgColor=couleur)))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:K{last}"
    return last


def synthese(wb, ao, last):
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = f"Appel d'offres · {ao.get('acheteur', '[à compléter]')} · {ao.get('reference', '')}"
    ws["A1"].font = Font(bold=True, size=14, color=TEXTE)
    ws["A2"] = ao.get("objet", "")
    infos = [("Date limite de remise", ao.get("date_limite")), ("Questions acceptées jusqu'au", ao.get("date_questions")),
             ("Critères de jugement", " · ".join(f"{c['critere']} {c['poids']} %" for c in ao.get("criteres", []))),
             ("Documents Triskell utilisés", " · ".join(ao.get("documents_triskell", [])) or "aucun : toutes les réponses sont à compléter"),
             ("Analyse du", ao.get("date_analyse")), ("Avertissement", ao.get("avertissement", ""))]
    for k, (lib, val) in enumerate(infos, start=4):
        ws.cell(k, 1, lib).font = Font(bold=True)
        ws.cell(k, 2, val or "[à compléter]").alignment = Alignment(wrap_text=True)
    r0 = 12
    ws.cell(r0 - 1, 1, "Couverture des exigences (se met à jour avec la matrice)").font = Font(bold=True, color=ORANGE)
    entete(ws, r0, ["Statut", "Obligatoires", "Souhaitées", "Total"])
    m = "'Matrice de conformité'"
    for k, statut in enumerate(STATUTS, start=r0 + 1):
        ws.cell(k, 1, statut)
        ws.cell(k, 2, f'=COUNTIFS({m}!$G$2:$G${last},A{k},{m}!$B$2:$B${last},"Obligatoire")')
        ws.cell(k, 3, f'=COUNTIFS({m}!$G$2:$G${last},A{k},{m}!$B$2:$B${last},"Souhaité")')
        ws.cell(k, 4, f'=COUNTIF({m}!$G$2:$G${last},A{k})')
    tot = r0 + len(STATUTS) + 1
    ws.cell(tot, 1, "Total").font = Font(bold=True)
    for j, col in ((2, "B"), (3, "C"), (4, "D")):
        ws.cell(tot, j, f"=SUM({col}{r0 + 1}:{col}{tot - 1})").font = Font(bold=True)
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "bar", "stacked", 100
    ch.title = "Exigences par statut"
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=r0, max_row=tot - 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=tot - 1))
    for s, col in zip(ch.series, (VERT, ORANGE)):
        s.graphicalProperties.solidFill = col
    ch.height, ch.width = 7, 13
    ws.add_chart(ch, f"F{r0 - 1}")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 70
    for col in "CD":
        ws.column_dimensions[col].width = 13


def questions(wb, qs):
    ws = wb.create_sheet("Questions acheteur")
    entete(ws, 1, ["N°", "Exigence", "Question à poser (par écrit)", "Pourquoi (impact sur la réponse)"], [6, 12, 70, 50])
    for i, q in enumerate(qs, start=2):
        for j, v in enumerate([i - 1, q.get("exigence"), q.get("question"), q.get("pourquoi")], start=1):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")


def go_no_go(wb, g):
    ws = wb.create_sheet("Grille go-no go")
    ws["A1"] = "Grille go / no go : notes de 0 à 3, à valider par la DC et la direction (le calcul n'est pas la décision)"
    ws["A1"].font = Font(bold=True, color=TEXTE)
    entete(ws, 3, ["Critère", "Poids", "Note (0-3)", "Justification (faits du dossier)"], [42, 9, 11, 80])
    crit = g.get("criteres", [])
    for i, c in enumerate(crit, start=4):
        for j, v in enumerate([c["critere"], c["poids"], c.get("note"), c.get("justification")], start=1):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    end = 3 + len(crit)
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="3", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C4:C{end}")
    s = end + 2
    ws.cell(s, 1, "Score pondéré (%)").font = Font(bold=True)
    ws.cell(s, 2, f"=IFERROR(ROUND(SUMPRODUCT(B4:B{end},C4:C{end})/(SUM(B4:B{end})*3)*100,0),\"\")")
    ws.cell(s + 1, 1, "Notes manquantes").font = Font(bold=True)
    ws.cell(s + 1, 2, f"=COUNTBLANK(C4:C{end})")
    ws.cell(s + 2, 1, "Lecture du score").font = Font(bold=True)
    ws.cell(s + 2, 2, f'=IF(B{s + 1}>0,"Incomplet : noter tous les critères",IF(B{s}>=70,"GO probable",IF(B{s}>=50,"À discuter","NO GO probable")))')
    ws.cell(s + 3, 1, "Un critère éliminatoire non levé = NO GO quel que soit le score.").font = Font(italic=True, color=ORANGE)
    e0 = s + 5
    entete(ws, e0, ["Critère éliminatoire", "Statut", "", "Commentaire"])
    for i, e in enumerate(g.get("eliminatoires", []), start=e0 + 1):
        for j, v in ((1, e["critere"]), (2, e.get("statut")), (4, e.get("commentaire"))):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("analyse")
    ap.add_argument("sortie")
    a = ap.parse_args()
    with open(a.analyse, encoding="utf-8") as f:
        data = json.load(f)
    wb = Workbook()
    last = matrice(wb, data["exigences"])
    synthese(wb, data.get("ao", {}), last)
    questions(wb, data.get("questions", []))
    go_no_go(wb, data.get("go_no_go", {}))
    wb.save(a.sortie)
    compte = {s: sum(1 for e in data["exigences"] if e["statut"] == s) for s in STATUTS}
    print(f"OK : {a.sortie} · {len(data['exigences'])} exigences · " + " · ".join(f"{k} {v}" for k, v in compte.items()))


if __name__ == "__main__":
    main()
