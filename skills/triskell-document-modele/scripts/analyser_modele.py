"""Relève la structure d'un document Word : paragraphes, styles, gras, tableaux, champs à remplir.

Usage :
    python analyser_modele.py modele.docx > structure.json
    python analyser_modele.py modele.docx --texte          # vue lisible, un paragraphe par ligne

Sert à deux choses :
1. construire la table de correspondance « section du modèle → contenu → source » ;
2. donner les numéros de paragraphe (`index`) utilisés par remplir_modele.py.

Champs à remplir détectés : [texte entre crochets], <texte entre chevrons>, XXX,
« NOM CLIENT », « A MODIFIER ». Le script ne modifie rien.
"""
import json
import re
import sys

import docx

MOTIF_CHAMP = re.compile(r"\[[^\[\]]{1,120}\]|<[^<>]{1,80}>|\bX{3,}\b|\bx{3,}\b|NOM CLIENT|A MODIFIER|À MODIFIER")


def champs(texte):
    return MOTIF_CHAMP.findall(texte or "")


def analyser(chemin):
    doc = docx.Document(chemin)
    paragraphes = []
    for i, p in enumerate(doc.paragraphs):
        runs = [{"texte": r.text, "gras": bool(r.bold), "italique": bool(r.italic)} for r in p.runs if r.text]
        paragraphes.append({
            "index": i,
            "style": p.style.name if p.style is not None else "",
            "texte": p.text,
            "runs": runs,
            "champs": champs(p.text),
            "vide": not p.text.strip(),
        })
    tableaux = []
    for t_i, t in enumerate(doc.tables):
        lignes = [[c.text for c in row.cells] for row in t.rows]
        tableaux.append({"index": t_i, "lignes": lignes,
                         "champs": sorted({c for l in lignes for cell in l for c in champs(cell)})})
    sec = doc.sections[0]
    entetes = [p.text for p in sec.header.paragraphs if p.text.strip()]
    pieds = [p.text for p in sec.footer.paragraphs if p.text.strip()]
    tous = [c for p in paragraphes for c in p["champs"]] + [c for t in tableaux for c in t["champs"]]
    return {
        "fichier": chemin,
        "paragraphes": paragraphes,
        "tableaux": tableaux,
        "en_tete_de_page": entetes,
        "pied_de_page": pieds,
        "champs_a_remplir": sorted(set(tous), key=tous.index),
        "styles_utilises": sorted({p["style"] for p in paragraphes}),
    }


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    resultat = analyser(sys.argv[1])
    if "--texte" in sys.argv:
        for p in resultat["paragraphes"]:
            if p["vide"]:
                continue
            gras = " [gras]" if p["runs"] and all(r["gras"] for r in p["runs"]) else ""
            print(f'{p["index"]:>3} | {p["style"]}{gras} | {p["texte"]!r}')
        for t in resultat["tableaux"]:
            print(f'Tableau {t["index"]} : {len(t["lignes"])} lignes, champs {t["champs"]}')
        print("Champs à remplir :", resultat["champs_a_remplir"])
        return
    json.dump(resultat, sys.stdout, ensure_ascii=False, indent=2)
    print()


if __name__ == "__main__":
    main()
