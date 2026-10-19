"""Produit le nouveau document en remplissant une COPIE du modèle Word (styles conservés).

Usage :
    python remplir_modele.py modele.docx plan.json sortie.docx

Le plan (JSON) dit quoi remplacer, quoi dupliquer, quoi supprimer. Les numéros de
paragraphe sont ceux d'analyser_modele.py (document d'origine). Format :

{
  "garder_a_partir_de": "texte du 1er paragraphe du modèle",   # facultatif : supprime ce qui précède
  "garder_jusqua": "texte du dernier paragraphe à garder",     # facultatif : supprime ce qui suit
  "dupliquer": [{"paragraphe": 20, "remplacements": [{"cherche": "[Point 3 — titre]", "remplace": "…"}]}],
  "remplacements": [{"cherche": "[Lieu]", "remplace": "Paris", "paragraphe": 13}],  # sans "paragraphe" = partout
  "supprimer_paragraphes": [19],
  "interdits": ["Hélianthe"]                                 # termes qui ne doivent plus apparaître
}

Le remplacement se fait run par run : gras, police et taille du modèle sont gardés.
Un « \\n » dans le texte devient un retour à la ligne. Tout texte contenant
« [à préciser » est surligné en jaune pour la relecture.

À la fin, le script contrôle la sortie : champs du modèle non remplis, termes
interdits restants. Code de sortie 2 si un terme interdit subsiste.
"""
import copy
import json
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX

sys.path.insert(0, str(Path(__file__).resolve().parent))
from analyser_modele import champs  # noqa: E402

MARQUE_A_PRECISER = "[à préciser"


def _tous_les_paragraphes(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            yield from part.paragraphs


def _surligner(run):
    if MARQUE_A_PRECISER in run.text:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def remplacer(paragraphe, cherche, remplace):
    """Remplace dans le paragraphe en gardant le format du run. Renvoie le nombre de remplacements."""
    n = 0
    for run in paragraphe.runs:
        if cherche in run.text:
            n += run.text.count(cherche)
            run.text = run.text.replace(cherche, remplace)
            _surligner(run)
    if n == 0 and cherche in paragraphe.text and paragraphe.runs:
        # Champ coupé entre plusieurs runs : on regroupe dans le premier run (format du premier run).
        texte = paragraphe.text.replace(cherche, remplace)
        paragraphe.runs[0].text = texte
        for run in paragraphe.runs[1:]:
            run.text = ""
        _surligner(paragraphe.runs[0])
        print(f"  ! « {cherche} » était coupé entre plusieurs runs : paragraphe regroupé, format du 1er run",
              file=sys.stderr)
        n = 1
    return n


def _hors_plage(doc, debut, fin):
    """Repère (avant toute modification) les éléments du corps à supprimer hors de la plage gardée."""
    elements = [e for e in doc.element.body.iterchildren() if not e.tag.endswith("sectPr")]
    textes = ["".join(t.text or "" for t in e.iter() if t.tag.endswith("}t")) for e in elements]

    def position(texte_cherche):
        for i, t in enumerate(textes):
            if t.strip().startswith(texte_cherche.strip()):
                return i
        sys.exit(f"Paragraphe introuvable : « {texte_cherche} »")

    i_debut = position(debut) if debut else 0
    i_fin = position(fin) if fin else len(elements) - 1
    return [e for i, e in enumerate(elements) if i < i_debut or i > i_fin]


def appliquer(plan, doc):
    paragraphes = list(doc.paragraphs)  # indices du document d'origine
    rapport = {"remplacements": 0, "non_trouves": []}
    a_supprimer = []
    if plan.get("garder_a_partir_de") or plan.get("garder_jusqua"):
        a_supprimer = _hors_plage(doc, plan.get("garder_a_partir_de"), plan.get("garder_jusqua"))

    for d in reversed(plan.get("dupliquer", [])):
        source = paragraphes[d["paragraphe"]]
        clone = copy.deepcopy(source._p)
        source._p.addnext(clone)
        nouveau = docx.text.paragraph.Paragraph(clone, source._parent)
        for r in d.get("remplacements", []):
            rapport["remplacements"] += remplacer(nouveau, r["cherche"], r["remplace"])

    for r in plan.get("remplacements", []):
        if "paragraphe" in r:
            cibles = [paragraphes[r["paragraphe"]]]
        else:
            cibles = list(_tous_les_paragraphes(doc))
        n = sum(remplacer(p, r["cherche"], r["remplace"]) for p in cibles)
        rapport["remplacements"] += n
        if n == 0:
            rapport["non_trouves"].append(r["cherche"])

    for i in sorted(plan.get("supprimer_paragraphes", []), reverse=True):
        p = paragraphes[i]._p
        p.getparent().remove(p)

    for e in a_supprimer:
        if e.getparent() is not None:
            e.getparent().remove(e)
    return rapport


def controler(doc, interdits):
    texte = "\n".join(p.text for p in _tous_les_paragraphes(doc))
    restants = [c for c in champs(texte) if not c.lower().startswith(MARQUE_A_PRECISER)]
    a_preciser = [c for c in champs(texte) if c.lower().startswith(MARQUE_A_PRECISER)]
    trouves = [t for t in interdits if t.lower() in texte.lower()]
    return {"champs_du_modele_restants": restants, "a_preciser": a_preciser, "interdits_restants": trouves}


def main():
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    modele, plan_json, sortie = sys.argv[1:]
    if Path(sortie).resolve() == Path(modele).resolve():
        sys.exit("La sortie doit être une copie : ne jamais écraser le modèle.")
    plan = json.loads(Path(plan_json).read_text(encoding="utf-8"))
    doc = docx.Document(modele)
    rapport = appliquer(plan, doc)
    doc.save(sortie)
    controle = controler(docx.Document(sortie), plan.get("interdits", []))
    print(json.dumps({"sortie": sortie, **rapport, **controle}, ensure_ascii=False, indent=2))
    if controle["interdits_restants"]:
        print("ATTENTION : des termes interdits subsistent, corriger le plan.", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
