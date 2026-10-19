"""Extrait les questions d'un questionnaire Excel et les entrées d'une base de réponses.

Usage :
    python extraire.py questions questionnaire.xlsx            > questions.json
    python extraire.py base base-reponses.docx|.xlsx|.md       > base.json

« questions » : repère, dans chaque feuille, la ligne d'en-tête (cellule contenant
« question ») et les colonnes utiles (n°, domaine, question, réponse, source, oui/non).
Une question = une ligne sous l'en-tête dont la cellule « question » est remplie.

« base » : découpe la base en entrées numérotées B1, B2… (un titre + son texte), pour
que chaque réponse du questionnaire puisse citer sa source exacte.

Le script ne modifie aucun fichier. Il n'interprète rien : la correspondance
question → réponse est faite par l'assistant, qui la relit.
"""
import json
import re
import sys
from pathlib import Path

ENTETES = {
    "question": r"question|exigence|requirement|contr[ôo]le|crit[èe]re",
    "reponse": r"r[ée]ponse|response|answer|commentaire du fournisseur",
    "source": r"source|justificatif|preuve|evidence|r[ée]f[ée]rence",
    "oui_non": r"oui\s*/\s*non|yes\s*/\s*no|conforme|statut fournisseur",
    "numero": r"^(n°|no|num|#|id|r[ée]f)\.?$",
    "domaine": r"domaine|th[èe]me|cat[ée]gorie|section|rubrique",
}
LIGNES_MAX_ENTETE = 20


def _norm(valeur):
    return str(valeur).strip().lower() if valeur is not None else ""


LONGUEUR_MAX_ENTETE = 60  # un titre de document n'est pas une cellule d'en-tête


def _trouver_entete(ws):
    """Ligne d'en-tête = celle qui reconnaît le plus de colonnes, dont « question »."""
    meilleure = (None, {})
    for row in ws.iter_rows(min_row=1, max_row=LIGNES_MAX_ENTETE):
        colonnes = {}
        for cell in row:
            texte = _norm(cell.value)
            if not texte or len(texte) > LONGUEUR_MAX_ENTETE:
                continue
            for cle, motif in ENTETES.items():
                if cle not in colonnes and re.search(motif, texte):
                    colonnes[cle] = cell.column
                    break
        if "question" in colonnes and len(colonnes) > len(meilleure[1]):
            meilleure = (row[0].row, colonnes)
    return meilleure


def extraire_questions(chemin):
    import openpyxl

    wb = openpyxl.load_workbook(chemin)
    feuilles = []
    for ws in wb.worksheets:
        ligne_entete, colonnes = _trouver_entete(ws)
        if not ligne_entete:
            feuilles.append({"feuille": ws.title, "entete": None, "questions": []})
            continue
        questions = []
        for r in range(ligne_entete + 1, ws.max_row + 1):
            question = ws.cell(r, colonnes["question"]).value
            if question is None or not str(question).strip():
                continue
            item = {"ligne": r, "question": str(question).strip()}
            for cle in ("numero", "domaine", "reponse"):
                if cle in colonnes:
                    v = ws.cell(r, colonnes[cle]).value
                    item[cle] = v if v is None or isinstance(v, (int, float)) else str(v).strip()
            questions.append(item)
        feuilles.append({
            "feuille": ws.title,
            "entete": ligne_entete,
            "colonnes": colonnes,
            "derniere_colonne": ws.max_column,
            "questions": questions,
        })
    total = sum(len(f["questions"]) for f in feuilles)
    return {"fichier": Path(chemin).name, "total_questions": total, "feuilles": feuilles}


def _base_docx(chemin):
    import docx

    doc = docx.Document(chemin)
    entetes, entrees, courante = [], [], None
    for p in doc.paragraphs:
        texte = p.text.strip()
        if not texte:
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading") or style.startswith("titre"):
            courante = {"titre": texte, "texte": ""}
            entrees.append(courante)
        elif style == "title":
            entetes.append(texte)
        elif courante is None:
            entetes.append(texte)
        else:
            courante["texte"] = (courante["texte"] + "\n" + texte).strip()
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0] and cells[-1]:
                entrees.append({"titre": cells[0], "texte": " | ".join(cells[1:])})
    return entetes, entrees


def _base_xlsx(chemin):
    import openpyxl

    wb = openpyxl.load_workbook(chemin, data_only=True)
    entrees = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if len(cells) >= 2:
                entrees.append({"titre": f"{ws.title} · {cells[0]}", "texte": " | ".join(cells[1:])})
    return [], entrees


def _base_texte(chemin):
    entetes, entrees, courante = [], [], None
    for ligne in Path(chemin).read_text(encoding="utf-8").splitlines():
        texte = ligne.strip()
        if not texte:
            continue
        if texte.startswith("#"):
            courante = {"titre": texte.lstrip("# ").strip(), "texte": ""}
            entrees.append(courante)
        elif courante is None:
            entetes.append(texte)
        else:
            courante["texte"] = (courante["texte"] + "\n" + texte).strip()
    return entetes, entrees


def extraire_base(chemin):
    suffixe = Path(chemin).suffix.lower()
    if suffixe == ".docx":
        entetes, entrees = _base_docx(chemin)
    elif suffixe in (".xlsx", ".xlsm"):
        entetes, entrees = _base_xlsx(chemin)
    else:
        entetes, entrees = _base_texte(chemin)
    entrees = [e for e in entrees if e["texte"]]
    return {
        "fichier": Path(chemin).name,
        "en_tete": entetes,
        "entrees": [{"id": f"B{i}", **e} for i, e in enumerate(entrees, start=1)],
    }


def main():
    if len(sys.argv) != 3 or sys.argv[1] not in ("questions", "base"):
        sys.exit(__doc__)
    mode, chemin = sys.argv[1], sys.argv[2]
    resultat = extraire_questions(chemin) if mode == "questions" else extraire_base(chemin)
    json.dump(resultat, sys.stdout, ensure_ascii=False, indent=2)
    print()


if __name__ == "__main__":
    main()
