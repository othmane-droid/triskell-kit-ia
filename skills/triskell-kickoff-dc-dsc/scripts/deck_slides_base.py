"""Slides de structure et de message : couverture, agenda, séparateur, fin, affirmation,
synthèse, puces, fiches, cartes, chiffres clés, citation, demande de décision."""
import os

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Cm

from deck_core import (C, FONT_BLACK, LAYOUT, TOKENS, ZONE, STATE, asset, check_forbidden, disc, fit_height,
                       panel, rect, shape, style_run, textbox, warn)

BOTTOM = ZONE["y"] + ZONE["h"]


def new_slide(prs, layout_key, title=None):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT[layout_key]])
    if layout_key == "contenu":
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
    if not title and layout_key == "contenu":
        slide.shapes.title._element.getparent().remove(slide.shapes.title._element)
        return slide
    if title is not None:
        check_forbidden(title)
        if len(title) > TOKENS["regles"]["titre_max_caracteres"]:
            warn(f"titre de {len(title)} caractères (max {TOKENS['regles']['titre_max_caracteres']})")
        slide.shapes.title.text = title.replace("**", "")
    return slide


def notes(slide, s):
    if s.get("notes"):
        slide.notes_slide.notes_text_frame.text = s["notes"]


def subtitle(slide, s, y=3.05):
    """Chapeau sous le titre (facultatif). Renvoie le y de départ du contenu."""
    if s.get("subtitle"):
        textbox(slide, ZONE["x"], y, ZONE["w"] - 4, 1.0, [{"text": s["subtitle"], "size": 16, "color": "texte_secondaire"}], min_size=12)
        return y + 1.35
    return ZONE["y"] + 0.2


def source(slide, s, required=False):
    if s.get("source"):
        textbox(slide, ZONE["x"], BOTTOM - 0.45, ZONE["w"] - 5, 0.7, [{"text": "Source : " + s["source"], "size": 10, "color": "texte_attenue"}], quiet=True)
    elif required:
        warn("chiffres sans source")


def content(prs, s):
    slide = new_slide(prs, "contenu", s["title"])
    return slide, subtitle(slide, s)


# ------------------------------------------------------------------ structure
def s_cover(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT["couverture"]])
    ph = {p.placeholder_format.idx: p for p in slide.placeholders}
    ph[1].text = s.get("kicker", "")
    ph[0].text = s["title"]
    if s.get("date_version"):
        p = ph[0].text_frame.add_paragraph()
        r = p.add_run()
        r.text = s["date_version"]
        style_run(r, 22, "texte_secondaire")
    if 10 in ph:
        ph[10]._element.getparent().remove(ph[10]._element)
    logo = s.get("client_logo")
    if logo and os.path.exists(logo):
        pic = slide.shapes.add_picture(logo, Cm(13.6), Cm(15.7), height=Cm(2.4))
        if pic.width > Cm(8):
            k = Cm(8) / pic.width
            pic.width, pic.height = Cm(8), int(pic.height * k)
    elif s.get("client_name"):
        textbox(slide, 2.95, 13.2, 14, 1.2, [{"text": "Pour " + s["client_name"], "size": 20, "bold": True, "color": "vert"}])
    if s.get("confidential", True):
        textbox(slide, 2.95, 17.3, 12, 0.8, [{"text": s.get("confidential_label", "Document confidentiel"), "size": 10, "color": "texte_attenue"}], quiet=True)
    notes(slide, s)


def s_agenda(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT["agenda"]])
    textbox(slide, 2.4, 1.3, 20, 3.4, [{"text": s.get("title", "Agenda"), "size": 66, "color": "orange", "font": FONT_BLACK, "bold": True}])
    items, cur = s["items"], s.get("current")
    if len(items) > 8:
        warn(f"{len(items)} points d'agenda (8 maximum)")
    y = 5.6
    step = min(1.55, 11.5 / max(1, len(items)))
    for i, it in enumerate(items, 1):
        on = cur is None or cur == i
        disc(slide, 2.4, y + 0.05, 1.05, "orange" if (cur == i) else ("vert" if on else "trait"), i, size=15)
        label = it if isinstance(it, str) else it["label"]
        textbox(slide, 3.9, y, 15.0, 1.2, [{"text": label, "size": 21, "bold": True, "color": "texte" if on else "texte_attenue"}], anchor="middle", min_size=14)
        if not isinstance(it, str) and it.get("time"):
            textbox(slide, 18.9, y, 3.0, 1.2, [{"text": it["time"], "size": 15, "color": "texte_secondaire", "align": "right"}], anchor="middle")
        y += step
    if s.get("qa", True):
        textbox(slide, 26.5, 9.0, 6.2, 3.0, [{"text": "Q&A", "size": 60, "color": "orange", "font": FONT_BLACK, "bold": True}])
    notes(slide, s)


def s_section(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT["separateur"]])
    label = s["title"] if not s.get("number") else f"{s['number']}. {s['title']}"
    check_forbidden(label)
    for ph in slide.placeholders:
        ph.text = label
    if s.get("teaser"):
        textbox(slide, 9.3, 10.9, 22, 2.0, [{"text": s["teaser"], "size": 18, "color": "blanc"}])
    notes(slide, s)


def s_closing(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT["fin"]])
    lines = []
    if s.get("contact"):
        lines.append({"text": s["contact"], "size": 17, "bold": True, "color": "texte"})
    for c in s.get("details", []):
        lines.append({"text": c, "size": 14, "color": "texte_secondaire"})
    if lines:
        textbox(slide, 2.95, 13.4, 22, 3.6, lines)
    notes(slide, s)


# ------------------------------------------------------------------ message
def s_statement(prs, s):
    """Une affirmation forte qui porte toute la slide. **mot** = mot en orange."""
    slide = new_slide(prs, "contenu", s.get("title", ""))
    rect(slide, ZONE["x"], 6.0, 0.28, 5.6, "orange")
    textbox(slide, ZONE["x"] + 1.2, 5.2, 26.5, 7.2, [{"text": s["text"], "size": 38, "bold": True, "color": "texte_titre", "font": FONT_BLACK}], anchor="middle", min_size=24)
    if s.get("support"):
        textbox(slide, ZONE["x"] + 1.2, 12.6, 24, 2.4, [{"text": s["support"], "size": 18, "color": "texte_secondaire"}], min_size=13)
    source(slide, s)
    notes(slide, s)


def s_summary(prs, s):
    """Synthèse décideur : 3 messages numérotés + la décision attendue."""
    slide, y = content(prs, s)
    msgs = s["messages"]
    left_w = 19.6 if s.get("decision") else ZONE["w"]
    step = min(3.6, (BOTTOM - y) / max(3, len(msgs)))
    for i, m in enumerate(msgs, 1):
        cy = y + (i - 1) * step
        disc(slide, ZONE["x"], cy + 0.1, 1.4, "orange" if i == 1 else "vert", i, size=18)
        textbox(slide, ZONE["x"] + 2.1, cy, left_w - 2.3, step - 0.3,
                [{"text": m["title"], "size": 19, "bold": True, "space_after": 4}] + ([{"text": m["text"], "size": 15, "color": "texte_secondaire"}] if m.get("text") else []), min_size=11)
    if s.get("decision"):
        d = s["decision"]
        x = ZONE["x"] + left_w + 0.8
        w = ZONE["w"] - left_w - 0.8
        dh = min(BOTTOM - y - 0.6, max(step * len(msgs) - 0.3, 6.5))
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, dh, d.get("color", "vert"), radius=0.05)
        paras = [{"text": d.get("label", "Décision attendue").upper(), "size": 12, "bold": True, "color": "blanc", "space_after": 10},
                 {"text": d["text"], "size": 21, "bold": True, "color": "blanc", "space_after": 12}]
        for o in d.get("options", []):
            paras.append({"text": o, "size": 14, "color": "blanc", "bullet": True, "bullet_color": "blanc" if d.get("color", "vert") != "vert" else "jaune"})
        textbox(slide, x + 0.7, y + 0.8, w - 1.4, dh - 1.4, paras, min_size=11)
    notes(slide, s)


def s_bullets(prs, s):
    slide, y = content(prs, s)
    groups = s.get("groups") or [{"heading": None, "items": s.get("items", [])}]
    cols = 1 if len(groups) == 1 else (3 if len(groups) == 3 else 2)
    gap = 0.8
    w = (ZONE["w"] - (cols - 1) * gap) / cols
    rows = -(-len(groups) // cols)
    avail = BOTTOM - y - 0.4
    for i, g in enumerate(groups):
        if len(g["items"]) > TOKENS["regles"]["puces_max_par_bloc"]:
            warn(f"{len(g['items'])} puces dans un bloc (max {TOKENS['regles']['puces_max_par_bloc']})")
        paras = [{"text": it, "size": 16, "bullet": True, "space_after": 8, "label": g.get("heading") or "point"} for it in g["items"]]
        h = avail / rows - (0.4 if rows > 1 else 0)
        if s.get("compact", False):
            h = fit_height(([{"text": g["heading"], "size": 17, "space_after": 8}] if g.get("heading") else []) + paras, w, 4.0, h)
        cx = ZONE["x"] + (i % cols) * (w + gap)
        cy = y + (i // cols) * (avail / rows)
        if s.get("plain") or (len(groups) == 1 and not g.get("heading")):
            textbox(slide, cx, cy, w, h, paras, min_size=11)
        else:
            panel(slide, cx, cy, w, h, g.get("heading"), paras, accent=g.get("accent", "orange"))
    notes(slide, s)


def s_fields(prs, s):
    """Fiches libellé → valeur. Valeur vide = [à compléter]. `inline` : « Libellé : valeur »."""
    slide, y = content(prs, s)
    cols = s["columns"]
    n = len(cols)
    gap = 0.8
    w = (ZONE["w"] - (n - 1) * gap) / n
    h = BOTTOM - y - 0.3
    for i, col in enumerate(cols):
        paras = []
        for f in col["fields"]:
            vals = f.get("value")
            if col.get("inline") and not isinstance(vals, list):
                paras.append({"prefix": f["label"] + " : ", "text": vals, "size": 15, "color": "texte_secondaire", "label": f["label"], "space_after": 8})
                continue
            paras.append({"text": f["label"], "size": 15, "bold": True, "space_after": 2})
            for v in (vals if isinstance(vals, list) else [vals]):
                paras.append({"text": v, "size": 14, "color": "texte_secondaire", "label": f["label"], "space_after": 9})
        panel(slide, ZONE["x"] + i * (w + gap), y, w, h, col.get("heading"), paras, accent=col.get("accent", "orange" if i % 2 == 0 else "vert"), min_size=9)
    notes(slide, s)


def s_cards(prs, s):
    slide, y = content(prs, s)
    cards = s["cards"]
    n = len(cards)
    if sum(1 for c in cards if c.get("highlight")) > 1:
        warn("plus d'une carte mise en avant")
    gap = 0.8
    w = (ZONE["w"] - (n - 1) * gap) / n
    h = min(BOTTOM - y - 0.3, s.get("card_height", 11.0))
    for i, c in enumerate(cards):
        cx = ZONE["x"] + i * (w + gap)
        hi = c.get("highlight")
        accent = "orange" if hi else "vert"
        rect(slide, cx, y, w, h, "fond_orange" if hi else "panneau", rounded=True)
        rect(slide, cx, y, w, 0.2, accent)
        disc(slide, cx + 0.6, y + 0.8, 1.5, accent, c.get("icon_text", i + 1), size=18)
        paras = []
        if c.get("kicker"):
            paras.append({"text": c["kicker"].upper(), "size": 12, "bold": True, "color": accent, "space_after": 6})
        paras.append({"text": c["title"], "size": 21, "bold": True, "space_after": 10})
        if c.get("text"):
            paras.append({"text": c["text"], "size": 15, "color": "texte_secondaire", "space_after": 8})
        for it in c.get("items", []):
            paras.append({"text": it, "size": 14, "bullet": True, "bullet_color": accent, "space_after": 5})
        textbox(slide, cx + 0.6, y + 2.7, w - 1.2, h - 3.1, paras, min_size=10)
        if hi and c.get("badge"):
            textbox(slide, cx + w - 4.6, y + 0.9, 4.0, 0.9, [{"text": c["badge"], "size": 11, "bold": True, "color": "blanc", "align": "center", "space_after": 0}], fill="orange", anchor="middle", pad=0.08)
    notes(slide, s)


def s_kpis(prs, s):
    slide, y = content(prs, s)
    k = s["kpis"]
    n = len(k)
    if sum(1 for x in k if x.get("highlight")) > 1:
        warn("plus d'un chiffre en orange")
    gap = 0.8
    w = (ZONE["w"] - (n - 1) * gap) / n
    h = min(9.0, BOTTOM - y - 1.4)
    for i, x in enumerate(k):
        cx = ZONE["x"] + i * (w + gap)
        hi = x.get("highlight")
        rect(slide, cx, y, w, h, "fond_orange" if hi else "panneau", rounded=True)
        rect(slide, cx, y, 0.2, h, "orange" if hi else "vert")
        textbox(slide, cx + 0.8, y + 0.8, w - 1.4, 3.2,
                [{"text": x["value"], "size": 54, "bold": True, "font": FONT_BLACK, "color": "orange" if hi else "texte_titre", "label": x.get("label", "chiffre")}], min_size=28)
        textbox(slide, cx + 0.8, y + 4.2, w - 1.4, h - 4.6,
                [{"text": x["label"], "size": 18, "bold": True, "space_after": 6}] + ([{"text": x["note"], "size": 14, "color": "texte_secondaire"}] if x.get("note") else []), min_size=11)
    if s.get("takeaway"):
        textbox(slide, ZONE["x"], y + h + 0.5, ZONE["w"], 2.0, [{"text": s["takeaway"], "size": 18, "bold": True, "color": "texte"}], min_size=13)
    source(slide, s, required=True)
    notes(slide, s)


def s_quote(prs, s):
    slide = new_slide(prs, "contenu", s.get("title", ""))
    textbox(slide, ZONE["x"] + 0.2, 2.6, 6, 7, [{"text": "“", "size": 150, "color": "orange", "font": FONT_BLACK, "bold": True}], quiet=True)
    textbox(slide, ZONE["x"] + 3.6, 5.4, 25.5, 6.4, [{"text": s["text"], "size": 28, "bold": True, "color": "texte_titre"}], anchor="middle", min_size=18)
    textbox(slide, ZONE["x"] + 3.6, 12.2, 24, 1.8, [{"text": s.get("author"), "size": 17, "color": "vert", "bold": True, "label": "auteur de la citation"},
                                                   {"text": s.get("role", ""), "size": 14, "color": "texte_secondaire", "allow_empty": True}])
    notes(slide, s)


def s_cta(prs, s):
    """Ce que nous vous demandons : la slide qui fait avancer la décision."""
    slide, y = content(prs, s)
    asks = s["asks"]
    lw = 20.0
    step = min(3.0, (BOTTOM - y - 0.4) / max(1, len(asks)))
    for i, a in enumerate(asks, 1):
        cy = y + (i - 1) * step
        rect(slide, ZONE["x"], cy, lw, step - 0.35, "panneau", rounded=True)
        disc(slide, ZONE["x"] + 0.5, cy + (step - 0.35 - 1.3) / 2, 1.3, "orange", i, size=16)
        textbox(slide, ZONE["x"] + 2.3, cy + 0.15, lw - 7.6, step - 0.65, [{"text": a["what"], "size": 17, "bold": True}], anchor="middle", min_size=11)
        textbox(slide, ZONE["x"] + lw - 5.1, cy + 0.15, 4.8, step - 0.65,
                [{"text": a.get("who"), "size": 14, "bold": True, "color": "vert", "label": "porteur", "space_after": 2},
                 {"text": a.get("when"), "size": 13, "color": "texte_secondaire", "label": "échéance"}], anchor="middle")
    x = ZONE["x"] + lw + 0.8
    w = ZONE["w"] - lw - 0.8
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, BOTTOM - y - 0.4, "vert", radius=0.05)
    paras = [{"text": s.get("next_label", "PROCHAIN RENDEZ-VOUS"), "size": 12, "bold": True, "color": "fond_doux", "space_after": 10},
             {"text": s.get("next"), "size": 22, "bold": True, "color": "blanc", "label": "prochain rendez-vous", "space_after": 14}]
    if s.get("contact"):
        paras.append({"text": s["contact"], "size": 14, "color": "blanc"})
    textbox(slide, x + 0.7, y + 0.8, w - 1.4, BOTTOM - y - 2.0, paras, min_size=11)
    notes(slide, s)


TYPES = {"cover": s_cover, "agenda": s_agenda, "section": s_section, "closing": s_closing, "statement": s_statement,
         "summary": s_summary, "bullets": s_bullets, "fields": s_fields, "cards": s_cards, "kpis": s_kpis,
         "quote": s_quote, "cta": s_cta}
