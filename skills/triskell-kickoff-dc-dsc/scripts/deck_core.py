"""Noyau du moteur Triskell : jetons, texte, ajustement automatique, formes.

Ne s'utilise pas seul : voir build_deck.py.
"""
import json
import math
import os
import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TOKENS = json.load(open(os.path.join(ROOT, "tokens", "tokens.json"), encoding="utf-8"))
C = {k: RGBColor.from_string(v.lstrip("#")) for k, v in TOKENS["couleurs"].items()}
FONT = TOKENS["polices"]["texte"]
FONT_BOLD = TOKENS["polices"]["titre"]
FONT_BLACK = TOKENS["polices"]["display"]
LAYOUT = TOKENS["dispositions"]
ZONE = TOKENS["zone_contenu_cm"]
MISSING = TOKENS["texte_manquant"]
FORBIDDEN = [s.lower() for s in TOKENS["consignes_interdites"]]
SERIES = [C[k] for k in TOKENS["series_graphiques"]]

STATE = {"n": 0, "title": "", "warnings": [], "missing": []}


def warn(msg):
    STATE["warnings"].append(f"slide {STATE['n']} « {STATE['title'][:40]} » : {msg}")


def asset(*parts):
    return os.path.join(ROOT, "assets", *parts)


def is_empty(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "xxx", "Xxx", "Xxxx", MISSING))


def note_missing(label):
    STATE["missing"].append({"slide": STATE["n"], "titre": STATE["title"], "champ": label})


def check_forbidden(text):
    low = (text or "").lower()
    for f in FORBIDDEN:
        if f in low:
            warn(f"consigne du modèle recopiée « {f} », à supprimer")


# ------------------------------------------------------------------ métrique du texte
def char_w(size):
    return size * 0.0352 * TOKENS["metrique"]["largeur_moyenne_caractere"]


def line_h(size):
    return size * 0.0352 * TOKENS["metrique"]["interligne"]


def lines_needed(text, size, width):
    per_line = max(1, int(width / char_w(size)))
    return sum(max(1, math.ceil(len(p) / per_line)) for p in str(text).split("\n"))


def strip_markup(t):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", str(t))


def height_needed(paras, width, ratio=1.0):
    total = 0.0
    for p in paras:
        sz = p.get("size", 14) * ratio
        txt = ("▪  " if p.get("bullet") else "") + p.get("prefix", "") + strip_markup(p["text"])
        total += lines_needed(txt, sz, width) * line_h(sz) + p.get("space_after", 4) * 0.0352
    return total


# ------------------------------------------------------------------ texte
def style_run(run, size, color="texte", bold=False, italic=False, font=None):
    f = run.font
    f.name = font or (FONT_BOLD if bold else FONT)
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = C[color]


def add_runs(par, text, size, color, bold, italic, font, accent="orange"):
    """« **mot** » dans un texte = mot en couleur d'accent et en gras (un seul accent par slide)."""
    parts = re.split(r"(\*\*.+?\*\*)", str(text))
    for part in parts:
        if not part:
            continue
        r = par.add_run()
        if part.startswith("**") and part.endswith("**"):
            r.text = part[2:-2]
            style_run(r, size, accent, True, italic, font)
        else:
            r.text = part
            style_run(r, size, color, bold, italic, font)


def normalize(paras):
    out = []
    for p in paras:
        p = {"text": p} if isinstance(p, str) else dict(p)
        if is_empty(p.get("text")):
            if p.get("allow_empty"):
                p["text"] = ""
            else:
                note_missing(p.get("label", "texte"))
                p.update(text=MISSING, color="orange", italic=True)
        check_forbidden(p["text"])
        out.append(p)
    return out


def textbox(slide, x, y, w, h, paras, size=14, min_size=10, anchor="top", align="left",
            fill=None, line=None, pad=None, quiet=False):
    """Zone de texte ajustée automatiquement. paras : chaînes ou dicts
    {text, size, bold, italic, color, bullet, bullet_color, prefix, space_after, align, font, accent}."""
    norm = normalize(paras)
    pad = (0.2 if h < 2 else 0.4) if (fill and pad is None) else (pad or 0)
    base = max([q.get("size", size) for q in norm] or [size])
    min_ratio = min(1.0, min_size / base)
    ratio = 1.0
    while ratio > min_ratio and height_needed([dict(q, size=q.get("size", size)) for q in norm], w - 2 * pad, ratio) > h - 2 * pad:
        ratio -= 0.03
    if ratio <= min_ratio and height_needed([dict(q, size=q.get("size", size)) for q in norm], w - 2 * pad, min_ratio) > h - 2 * pad and not quiet:
        warn(f"texte trop long pour sa zone (« {strip_markup(norm[0]['text'])[:45]} »), à raccourcir")
    ratio = max(ratio, min_ratio)
    if fill:
        shp = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line=line, radius=0.06)
        tf = shp.text_frame
    else:
        tf = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h)).text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Cm(pad))
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for i, q in enumerate(norm):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[q.get("align", align)]
        par.space_after = Pt(q.get("space_after", 4))
        sz = q.get("size", size) * ratio
        if q.get("bullet"):
            r0 = par.add_run()
            r0.text = "▪  "
            style_run(r0, sz, q.get("bullet_color", "orange"), True)
        if q.get("prefix"):
            rp = par.add_run()
            rp.text = q["prefix"]
            style_run(rp, sz, q.get("prefix_color", "texte"), True)
        add_runs(par, q["text"], sz, q.get("color", "texte"), q.get("bold", False), q.get("italic", False), q.get("font"), q.get("accent", "orange"))
    return tf


# ------------------------------------------------------------------ formes
def shape(slide, kind, x, y, w, h, fill, line=None, radius=None, line_w=1.0):
    shp = slide.shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h))
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = C[fill]
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = C[line]
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def rect(slide, x, y, w, h, color, rounded=False):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h, color, radius=0.08 if rounded else None)


def line(slide, x1, y1, x2, y2, color="trait", width=1.25):
    conn = slide.shapes.add_connector(1, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    conn.line.color.rgb = C[color]
    conn.line.width = Pt(width)
    return conn


def disc(slide, x, y, d, color, text=None, size=16, text_color="blanc"):
    shp = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, color)
    if text is not None:
        tf = shp.text_frame
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", Cm(0))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(text)
        style_run(r, size, text_color, True, font=FONT_BLACK)
    return shp


def panel(slide, x, y, w, h, heading=None, paras=(), accent="orange", fill="panneau", size=15, min_size=10, heading_size=17):
    """Panneau clair avec liseré d'accent en haut et intertitre : la brique de base des slides de contenu."""
    rect(slide, x, y, w, h, fill, rounded=True)
    rect(slide, x + 0.5, y, 2.2, 0.16, accent)
    items = ([{"text": heading, "size": heading_size, "bold": True, "color": accent if accent != "vert_clair" else "vert", "space_after": 8}] if heading else []) + list(paras)
    textbox(slide, x + 0.55, y + 0.55, w - 1.1, h - 0.9, items, size=size, min_size=min_size)


def fit_height(paras, w, lo, hi):
    """Hauteur de panneau ajustée au contenu, bornée."""
    return max(lo, min(hi, height_needed(normalize_preview(paras), w - 1.1) + 1.3))


def normalize_preview(paras):
    return [({"text": p} if isinstance(p, str) else dict(p, text=p.get("text") or MISSING)) for p in paras]
